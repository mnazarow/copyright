# -*- coding: utf-8 -*-
"""Аналитика отчетов и тренды.

Для каждого готового задания один раз считается сводка (stats.json рядом с state.json, пересчитывается при
изменении state.json). Страницы «Аналитика» и «Тренды» агрегируют эти сводки с фильтрами.
"""
import collections
import datetime as dt
import json
import math
import os
import re
import statistics
import threading

from . import core
from .pipeline.common import AUTHOR, TYPES
from .pipeline.exports import safe

STATS_VERSION = 7

# правило → (название, источник, ключ настройки)
RULES = {
    "range": ("Диапазон чисел: тире вместо дефиса", "typo", "passes"),
    "dot": ("Пробел после точки", "typo", "passes"),
    "digitword": ("Пробел между числом и словом", "typo", "passes"),
    "dspace": ("Двойной пробел", "typo", "passes"),
    "degree": ("Знак градуса", "typo", "passes"),
    "quotes": ("Вид кавычек", "typo", "quote_style"),
    "inch": ("Дюймы", "typo", "passes"),
    "equals": ("Отбивка знака =", "typo", "passes"),
    "unit": ("Единица слитно с числом", "typo", "passes"),
    "times": ("Знак × без отбивки", "typo", "passes"),
    "mwc": ("Метры водяного столба", "typo", "passes"),
    "thousands": ("Разряды в числах", "typo", "thousands"),
    "decimal": ("Десятичный разделитель", "typo", "decimal_separator"),
    "std": ("Обозначения стандартов", "data", "check_standards"),
    "pct": ("Пробел перед %", "typo", "percent_space"),
    "hanging": ("Висячие предлоги и союзы", "typo", "hanging_check"),
    "mixed": ("Смешение алфавитов в слове", "codes", "check_codes"),
    "latin": ("Латинская буква в русском тексте", "codes", "check_codes"),
    "glossary": ("Глоссарий компании", "house", "glossary"),
    "forbidden": ("Запрещенные слова", "house", "forbidden_words"),
    "brand": ("Написание бренда", "house", "brand_case"),
    "table": ("Данные таблиц", "data", "check_tables"),
    "toc": ("Содержание и номера страниц", "data", "check_toc"),
    "ref": ("Ссылки «стр. N»", "data", "check_refs"),
    "figures": ("Нумерация рисунков и таблиц", "data", "check_figures"),
    "phone": ("Формат телефонов", "data", "check_phone_format"),
    "requisites": ("Реквизиты компании", "data", "requisites"),
    "spell": ("Слово вне словаря", "spell", "hunspell_enabled"),
    "lt": ("LanguageTool", "lt", "lt_enabled"),
    "ai": ("Языковая модель", "ai", "ai_provider"),
    "batch": ("Сравнение пакета", "batch", "batch_compare"),
    "candidate": ("Неподтвержденный кандидат", "cand", "unresolved_candidates"),
}
SOURCES = {
    "typo": "Типографские правила",
    "codes": "Коды и алфавиты",
    "house": "Стандарты компании",
    "data": "Проверка данных",
    "spell": "Словарь Hunspell",
    "lt": "LanguageTool",
    "ai": "Языковая модель",
    "batch": "Сравнение пакета",
    "cand": "Кандидаты extract.py",
    "other": "Прочее",
}
STATUS = {"accepted": "Принято", "rejected": "Отклонено", "new": "Не просмотрено"}
PREFLIGHT = [("fonts", "Шрифты не встроены"), ("low_dpi", "Изображения ниже порога dpi"),
             ("scans", "Страницы без текстового слоя"), ("placeholders", "Серые заглушки"),
             ("empty", "Пустые страницы"), ("sizes", "Разные форматы страниц"),
             ("curves", "Текст в кривых"), ("style", "Нарушения стандарта оформления"), ("products", "Расхождения со справочником продукции"),
             ("images", "Одинаковые изображения у разных изделий")]


_learn_cache = {}
DECISION_EPOCH = [0]          # увеличивается при каждом решении корректора (✓ ✕, массовая отметка, удаление задания)
_epoch_lock = threading.Lock()


def decisions_changed():
    with _epoch_lock:                                # «+=» из нескольких потоков терял увеличения
        DECISION_EPOCH[0] += 1


def learn_key(rule, bad):
    """Ключ «правило + фрагмент»: фрагмент обрезается так же, как в сводках (80 символов)."""
    from .pipeline.common import norm
    return (rule or "", norm((bad or "")[:80]).lower())


def learned_rejections(threshold, exclude_job=None):
    """Пары (правило, фрагмент), отклоненные корректором не меньше чем в threshold разных документах и ни разу не
    принятые. Замечания, скрытые автоматически, не учитываются (иначе правило само себя усиливало бы).
    Результат кешируется до следующего решения корректора."""
    with core.db() as c:
        done = c.execute("SELECT COUNT(*) FROM jobs WHERE status='done'").fetchone()[0]
    ck = (DECISION_EPOCH[0], threshold, done)
    res = _learn_cache.get(ck)                       # .get: другой поток может очистить кэш между проверкой и чтением
    if res is None:
        acc, rej = collections.Counter(), collections.defaultdict(set)
        for s in all_stats():
            for x in s["findings"]:
                if x.get("auto"):
                    continue
                k = learn_key(x["rule"], x["bad"])
                if x["status"] == "accepted":
                    acc[k] += 1
                elif x["status"] == "rejected":
                    rej[k].add(s["id"])
        res = {k: docs for k, docs in rej.items() if len(docs) >= threshold and not acc.get(k) and k[1]}
        _learn_cache.clear()
        _learn_cache[ck] = res
    return {k: len(docs - {exclude_job}) for k, docs in res.items() if len(docs - {exclude_job}) >= threshold}


def rule_info(rule):
    lab, src, key = RULES.get(rule or "", (rule or "—", "other", ""))
    return {"rule": rule or "", "label": lab, "source": src, "source_label": SOURCES[src], "setting": key}


# ------------------------------------------------------------------ сводка одного задания
def family_of(name):
    """Семейство документа: имя без номеров версий, дат и служебных слов — версии одного каталога."""
    s = (name or "")[:200].lower().replace("_", " ").replace(".", " ")
    # «\d+(\s*\d+)*» перебирал все разбиения длинной цепочки цифр: имя файла вроде «прайс ред 2024100912345678901234x»
    # занимало процесс на минуты и часы (регулярное выражение держит GIL — замирал весь сервис, вместе с /api/health)
    s = re.sub(r"\b(v|в|вер|версия|ver|rev|ред|редакция)\s*\d+(?:\s+\d+)*\b", " ", s)
    s = re.sub(r"\(\d+\)|\b\d{1,2}[ .-]\d{1,2}[ .-]\d{2,4}\b|\b20\d\d\b|\bfinal\b|\bитог\w*|\bисправл\w*|\bновая\b|\bкопия\b", " ", s)
    s = re.sub(r"\bстраницы\b", " ", s)
    return re.sub(r"\s+", " ", s).strip(" -") or (name or "").lower()


def _stats_from_state(state):
    pages = int(state.get("pages") or 0) or 1
    findings = []
    for f in state.get("findings", []):
        ri = rule_info(f.get("rule"))
        findings.append({
            "id": f.get("id"), "page": int(f.get("page") or 1), "type": f.get("type", ""), "rule": ri["rule"],
            "source": ri["source"], "status": f.get("status", "new"), "systemic": int(f.get("systemic") or 0),
            "bad": (f.get("bad") or "")[:80], "good": (f.get("good") or "")[:80], "edited": bool(f.get("edited")),
            "auto": f.get("auto") or "",
        })
    v = state.get("verify") or {}
    pf = state.get("preflight") or {}
    pre = {
        "fonts": len(pf.get("fonts_not_embedded") or []),
        "low_dpi": len({x.get("page") for x in pf.get("low_dpi") or []}),
        "scans": len(pf.get("scans") or []),
        "placeholders": len(pf.get("placeholders") or []),
        "empty": len(pf.get("empty") or []),
        "sizes": 1 if len(pf.get("page_sizes") or {}) > 1 else 0,
        "curves": len(pf.get("curves") or []),
        "style": len((state.get("style") or {}).get("issues") or []) if (state.get("style") or {}).get("configured") else 0,
        "products": len((state.get("products") or {}).get("issues") or []),
        "images": len((state.get("images") or {}).get("duplicates") or []),
    }
    return {
        "_v": STATS_VERSION, "pages": pages, "findings": findings,
        "verify_ok": bool(v.get("ok")), "verify_passed": v.get("passed", 0), "verify_total": v.get("total", 0),
        "verify_fail": [c["name"] for c in v.get("checks", []) if not c.get("ok")],
        "preflight": pre, "auto_rejected": len(state.get("rejected") or []),
        "reviewed_cats": len(state.get("reviewed") or []),
    }


_stats_mem = collections.OrderedDict()   # номер задания → (ключ свежести, сводка): не читать диск на каждый запрос
_stats_mem_lock = threading.Lock()
_stats_mem_size = [0]                    # замечаний в кеше (≈1 КБ памяти каждое)
STATS_MEM_MAX = 200_000                  # больше — самые давно использованные сводки вытесняются (читаются с диска)


def _cache_put(jid, key, st):
    if os.environ.get("KORREKTOR_PARENT_PID"):   # процесс задания: живет недолго, кеш не нужен
        return
    n = len(st.get("findings") or ())
    with _stats_mem_lock:
        old = _stats_mem.pop(jid, None)
        if old:
            _stats_mem_size[0] -= len(old[1].get("findings") or ())
        _stats_mem[jid] = (key, st)
        _stats_mem_size[0] += n
        while _stats_mem_size[0] > STATS_MEM_MAX and len(_stats_mem) > 1:
            _, (_, ev) = _stats_mem.popitem(last=False)
            _stats_mem_size[0] -= len(ev.get("findings") or ())


def _cache_drop(jid):
    old = _stats_mem.pop(jid, None)
    if old:
        _stats_mem_size[0] -= len(old[1].get("findings") or ())
_JOB_KEYS = ("name", "batch", "created", "started", "finished", "author", "department", "owner", "doc_type")


def _mtime(p):
    try:
        return os.path.getmtime(p)
    except OSError:
        return None


def job_stats(job):
    """Сводка задания (кеш в памяти и в stats.json). None — если задание не завершено. Возвращается копия
    верхнего уровня; списки замечаний общие с кешем — их не изменять."""
    if job.get("status") != "done":
        return None
    d = core.job_dir(job["id"])
    mt = _mtime(os.path.join(d, "state.json"))
    if mt is None:
        return None
    key = (mt, _mtime(os.path.join(d, "timing.json"))) + tuple(job.get(k) for k in _JOB_KEYS)
    with _stats_mem_lock:
        hit = _stats_mem.get(job["id"])
        if hit and hit[0] == key:
            _stats_mem.move_to_end(job["id"])
    if hit and hit[0] == key:
        return dict(hit[1])
    st = _job_stats_disk(job, d, mt)
    if st is not None:
        _cache_put(job["id"], key, st)
        st = dict(st)
    return st


def _job_stats_disk(job, d, mt):
    sp = os.path.join(d, "state.json")
    cp = os.path.join(d, "stats.json")
    st = None
    if os.path.exists(cp):
        try:
            st = json.load(open(cp, encoding="utf-8"))
            if st.get("_v") != STATS_VERSION or st.get("_mtime") != mt:
                st = None
        except Exception:  # noqa
            st = None
    if st is None:
        try:
            with open(sp, encoding="utf-8") as fh:
                state = json.load(fh)
        except Exception:  # noqa
            return None
        st = _stats_from_state(state)
        st["_mtime"] = mt
        try:
            core.write_json_atomic(cp, st)
        except OSError:
            pass
    created = job.get("created") or ""
    start = job.get("started") or created
    dur = None
    if job.get("finished") and start:
        try:
            dur = max(0.0, (dt.datetime.fromisoformat(job["finished"]) - dt.datetime.fromisoformat(start)).total_seconds())
        except ValueError:
            dur = None
    tp = os.path.join(d, "timing.json")
    st["stages"] = []
    if os.path.exists(tp):
        try:
            st["stages"] = json.load(open(tp, encoding="utf-8"))
        except Exception:  # noqa
            pass
    st.update(author=job.get("author") or "", department=job.get("department") or "", owner=job.get("owner") or "",
              doc_type=job.get("doc_type") or "")
    st.update(id=job["id"], name=job["name"], batch=job.get("batch") or "", created=created, duration=dur,
              family=family_of(job["name"]))
    return st


def all_stats():
    out = []
    ids = set()
    for j in core.list_jobs(100000):
        ids.add(j["id"])
        s = job_stats(j)
        if s:
            out.append(s)
    with _stats_mem_lock:                            # удаленные задания — из кеша
        for jid in [k for k in _stats_mem if k not in ids]:
            _cache_drop(jid)
    out.sort(key=lambda s: (s["created"], s["id"]))
    return out


# расчеты по всем документам: одновременно не больше двух — быстрые смены фильтров на странице аналитики
# не запускают десятки одинаковых расчетов, занимающих процессор и память
_compute = threading.BoundedSemaphore(2)


def _limited(fn):
    import functools

    @functools.wraps(fn)
    def wrap(*a, **kw):
        with _compute:
            return fn(*a, **kw)
    return wrap


def failed_jobs(date_from=None, date_to=None, owner=None):
    n = 0
    for j in core.list_jobs(100000, owner=owner):
        if j["status"] == "failed" and _in_range(j.get("created", ""), date_from, date_to):
            n += 1
    return n


# ------------------------------------------------------------------ фильтры
def _d(s):
    try:
        d = dt.date.fromisoformat(str(s)[:10]) if s else None
    except ValueError:
        return None
    if d and not (dt.date(1990, 1, 1) <= d <= dt.date(2200, 1, 1)):
        return None                                    # «0001-01-01», «9999-12-31» — не даты документов
    return d


def _in_range(created, date_from, date_to):
    day = _d(created)
    if day is None:
        return False
    return (not date_from or day >= date_from) and (not date_to or day <= date_to)


def parse_filters(q):
    """Фильтры из параметров запроса. period: 7, 30, 90, 180, 365, all или custom (from/to)."""
    today = dt.date.today()
    period = q.get("period") or "all"
    df, dt_ = _d(q.get("from")), _d(q.get("to"))
    if period not in ("custom", "all"):
        try:
            df, dt_ = today - dt.timedelta(days=min(36600, max(1, int(period))) - 1), today
        except (ValueError, OverflowError):
            df = dt_ = None
    elif period == "all":
        df = dt_ = None
    if df and dt_ and df > dt_:
        df, dt_ = dt_, df
    return {
        "period": period, "from": df, "to": dt_, "q": (q.get("q") or "").strip().lower(),
        "batch": q.get("batch") or "", "family": q.get("family") or "", "job": q.get("job") or "",
        "type": q.get("type") or "", "source": q.get("source") or "", "status": q.get("status") or "",
        "rule": q.get("rule") or "", "author": q.get("author") or "", "department": q.get("department") or "",
        "owner": q.get("owner"), "doc_type": q.get("doc_type") or "",
    }


def _visible(scope, s):
    """scope — логин (str) или users.Scope."""
    if isinstance(scope, str):
        return s.get("owner") == scope
    return scope.ok(s.get("owner"), s.get("department"))


def filter_docs(stats, f):
    out = []
    for s in stats:
        if not _in_range(s["created"], f["from"], f["to"]):
            continue
        if f["q"] and f["q"] not in s["name"].lower():
            continue
        if f["batch"] and s["batch"] != f["batch"]:
            continue
        if f["family"] and s["family"] != f["family"]:
            continue
        if f.get("owner") is not None and not _visible(f["owner"], s):
            continue
        if f.get("doc_type") and s.get("doc_type") != f["doc_type"]:
            continue
        if f.get("author") and s.get("author") != f["author"]:
            continue
        if f.get("department") and s.get("department") != f["department"]:
            continue
        if f["job"] and str(s["id"]) != str(f["job"]):
            continue
        out.append(s)
    return out


def finding_ok(x, f):
    return ((not f["type"] or x["type"] == f["type"]) and (not f["source"] or x["source"] == f["source"])
            and (not f["status"] or x["status"] == f["status"]) and (not f["rule"] or x["rule"] == f["rule"]))


def _pct(a, b):
    return round(100.0 * a / b, 1) if b else None


def _r(x, n=2):
    return round(x, n) if x is not None else None


# ------------------------------------------------------------------ показатели
def metrics(docs, f=None):
    """Базовые показатели набора документов (используются и в аналитике, и в трендах)."""
    f = f or {"type": "", "source": "", "status": "", "rule": ""}
    pages = sum(s["pages"] for s in docs)
    fs = [x for s in docs for x in s["findings"] if finding_ok(x, f)]
    n = len(fs)
    occ = sum(max(1, x["systemic"]) for x in fs)
    acc = sum(1 for x in fs if x["status"] == "accepted")
    rej = sum(1 for x in fs if x["status"] == "rejected" and not x.get("auto"))
    auto = sum(1 for x in fs if x.get("auto"))
    durs = [s["duration"] for s in docs if s["duration"] is not None]
    return {
        "docs": len(docs), "pages": pages, "findings": n, "occurrences": occ,
        "density": _r(10.0 * n / pages) if pages else None,
        "per_doc": _r(n / len(docs), 1) if docs else None,
        "author_share": _pct(sum(1 for x in fs if x["type"] == AUTHOR), n),
        "ai_share": _pct(sum(1 for x in fs if x["source"] == "ai"), n),
        "accepted": acc, "rejected": rej, "new": n - acc - rej - auto, "auto_hidden": auto,
        "precision": _pct(acc, acc + rej),
        "reviewed_share": _pct(acc + rej, n),
        "edited": sum(1 for x in fs if x.get("edited")),
        "verify_rate": _pct(sum(1 for s in docs if s["verify_ok"]), len(docs)),
        "avg_duration": _r(statistics.mean(durs), 1) if durs else None,
        "median_duration": _r(statistics.median(durs), 1) if durs else None,
        "sec_per_page": _r(sum(durs) / sum(s["pages"] for s in docs if s["duration"] is not None), 2) if durs else None,
        "auto_rejected": sum(s["auto_rejected"] for s in docs),
        "systemic": sum(1 for x in fs if x["systemic"]),
    }


@_limited
def analytics(q):
    f = parse_filters(q)
    stats = all_stats()
    docs = filter_docs(stats, f)
    fs = [(s, x) for s in docs for x in s["findings"] if finding_ok(x, f)]
    m = metrics(docs, f)
    m["failed"] = failed_jobs(f["from"], f["to"], f.get("owner"))

    by_type = collections.Counter(x["type"] for _, x in fs)
    by_source = collections.Counter(x["source"] for _, x in fs)
    by_status = collections.Counter(x["status"] for _, x in fs)

    # точность по правилам
    rules = collections.defaultdict(lambda: {"accepted": 0, "rejected": 0, "new": 0, "docs": set()})
    for s, x in fs:
        r = rules[x["rule"]]
        r[x["status"] if x["status"] in ("accepted", "rejected") and not x.get("auto") else "new"] += 1
        r["docs"].add(s["id"])
    rule_rows = []
    for rule, r in rules.items():
        ri = rule_info(rule)
        tot = r["accepted"] + r["rejected"] + r["new"]
        rule_rows.append(dict(ri, n=tot, accepted=r["accepted"], rejected=r["rejected"], new=r["new"],
                              docs=len(r["docs"]), precision=_pct(r["accepted"], r["accepted"] + r["rejected"])))
    rule_rows.sort(key=lambda r: -r["n"])

    # плотность по документам
    dens = []
    for s in docs:
        n = sum(1 for x in s["findings"] if finding_ok(x, f))
        dens.append({"id": s["id"], "name": s["name"], "pages": s["pages"], "findings": n,
                     "density": _r(10.0 * n / s["pages"]), "verify_ok": s["verify_ok"], "created": s["created"][:10]})
    dens.sort(key=lambda r: -(r["density"] or 0))

    # документ × тип
    heat_rows = []
    for s in sorted(docs, key=lambda s: -len(s["findings"]))[:30]:
        c = collections.Counter(x["type"] for x in s["findings"] if finding_ok(x, f))
        heat_rows.append({"id": s["id"], "name": s["name"], "values": [c.get(t, 0) for t in TYPES]})

    # положение в документе (десятые доли объема)
    pos = [0] * 10
    for s, x in fs:
        pos[min(9, int((x["page"] - 1) * 10 / max(1, s["pages"])))] += 1

    # замечаний на документ — гистограмма
    per_doc = [r["findings"] for r in dens]
    edges = [0, 5, 10, 20, 30, 50, 75, 100, 10 ** 9]
    labels = ["0–4", "5–9", "10–19", "20–29", "30–49", "50–74", "75–99", "100+"]
    hist = [sum(1 for v in per_doc if edges[i] <= v < edges[i + 1]) for i in range(len(labels))]

    # объем документов
    pedges = [1, 5, 13, 25, 51, 101, 201, 10 ** 9]
    plabels = ["1–4", "5–12", "13–24", "25–50", "51–100", "101–200", "200+"]
    phist = [sum(1 for s in docs if pedges[i] <= s["pages"] < pedges[i + 1]) for i in range(len(plabels))]

    # время обработки
    scatter = [{"x": s["pages"], "y": _r(s["duration"], 1), "id": s["id"], "name": s["name"]}
               for s in docs if s["duration"] is not None]

    # предпечатная проверка: сколько документов затронуто
    pre = [{"key": k, "label": lab, "docs": sum(1 for s in docs if s["preflight"].get(k)),
            "items": sum(s["preflight"].get(k, 0) for s in docs)} for k, lab in PREFLIGHT]

    # этапы обработки: среднее время
    stg = collections.defaultdict(list)
    for s in docs:
        for r in s.get("stages") or []:
            stg[r["step"]].append(r["sec"])
    stages = [{"label": k, "n": _r(statistics.mean(v), 1), "docs": len(v)} for k, v in stg.items()]

    # машинная проверка
    vf = collections.Counter(n for s in docs for n in s["verify_fail"])

    # повторяющиеся ошибки (одна и та же пара «ошибка → исправление» в разных документах)
    rec = collections.defaultdict(lambda: {"n": 0, "docs": set(), "type": "", "rule": "", "good": ""})
    for s, x in fs:
        if not x["bad"] or x["source"] == "batch":
            continue
        k = (x["bad"].lower(), x["good"].lower() if x["good"] not in ("—", "") else "")
        r = rec[k]
        r["n"] += max(1, x["systemic"])
        r["docs"].add(s["id"])
        r["type"], r["rule"], r["good"], r["bad"] = x["type"], x["rule"], x["good"], x["bad"]
    recurring = sorted(({"bad": r["bad"], "good": r["good"], "type": r["type"], "rule": r["rule"],
                         "rule_label": rule_info(r["rule"])["label"], "n": r["n"], "docs": len(r["docs"]),
                         "doc_ids": sorted(r["docs"])[:20]}
                        for r in rec.values()), key=lambda r: (-r["docs"], -r["n"]))[:50]

    # слова вне словаря
    st = core.load_settings()
    wl = set((st.get("whitelist") or "").split())
    wl_low = {w.lower() for w in wl}
    unk = collections.defaultdict(lambda: {"n": 0, "docs": set(), "accepted": 0, "rejected": 0})
    for s, x in fs:
        if x["rule"] != "spell" or not x["bad"]:
            continue
        w = x["bad"].strip(" .,;:«»\"()")
        u = unk[w]
        u["n"] += 1
        u["docs"].add(s["id"])
        if x["status"] in ("accepted", "rejected"):
            u[x["status"]] += 1
    unknown = sorted(({"word": w, "n": u["n"], "docs": len(u["docs"]), "accepted": u["accepted"],
                       "rejected": u["rejected"], "whitelisted": w.lower() in wl_low}
                      for w, u in unk.items()), key=lambda r: (-r["docs"], -r["n"], r["word"]))[:80]

    # системные и единичные
    systemic = {"single": sum(1 for _, x in fs if not x["systemic"]),
                "systemic": sum(1 for _, x in fs if x["systemic"]),
                "systemic_occ": sum(x["systemic"] for _, x in fs if x["systemic"])}

    # страницы с наибольшим числом замечаний
    pages_top = collections.Counter((s["id"], s["name"], x["page"]) for s, x in fs).most_common(15)

    res = {
        "filters": {k: (v.isoformat() if isinstance(v, dt.date) else v) for k, v in f.items() if k != "owner"},
        "options": options([s for s in stats if f.get("owner") is None or _visible(f["owner"], s)]),
        "kpi": m,
        "by_type": [{"key": t, "label": t, "n": by_type.get(t, 0)} for t in TYPES],
        "by_source": [{"key": k, "label": SOURCES[k], "n": n} for k, n in by_source.most_common()],
        "by_status": [{"key": k, "label": lab, "n": by_status.get(k, 0)} for k, lab in STATUS.items()],
        "top_rules": [{"key": r["rule"], "label": r["label"], "n": r["n"], "source": r["source"]} for r in rule_rows[:15]],
        "rules": rule_rows,
        "density": dens,
        "heatmap": {"cols": TYPES, "rows": heat_rows},
        "position": [{"label": f"{i * 10}–{i * 10 + 10} %", "n": v} for i, v in enumerate(pos)],
        "per_doc_hist": [{"label": lab, "n": v} for lab, v in zip(labels, hist)],
        "pages_hist": [{"label": lab, "n": v} for lab, v in zip(plabels, phist)],
        "duration": scatter,
        "preflight": pre,
        "verify_fail": [{"label": k, "n": v} for k, v in vf.most_common()],
        "stages": stages,
        "recurring": recurring,
        "unknown": unknown,
        "systemic": systemic,
        "pages_top": [{"id": i, "name": nm, "page": p, "n": n} for (i, nm, p), n in pages_top],
        "by_author": group_by(docs, f, "author"),
        "by_department": group_by(docs, f, "department"),
        "by_doc_type": group_by(docs, f, "doc_type"),
        "goals": goals(),
    }
    g = res["goals"]
    st_ = core.load_settings()
    by_id = {s["id"]: s for s in docs}
    m["over_goal"] = 0
    for r in dens:
        gl = density_goal(by_id[r["id"]].get("doc_type"), st_)
        r["goal"] = gl
        if gl and (r["density"] or 0) > gl:
            m["over_goal"] += 1
    res["goal_state"] = {k: goal_state(k, m.get(k), g) for k in g}
    res["recommendations"] = recommendations(res, core.load_settings())
    return res


def plural(n, one, few, many):
    n10, n100 = n % 10, n % 100
    w = one if n10 == 1 and n100 != 11 else few if 2 <= n10 <= 4 and not 12 <= n100 <= 14 else many
    return f"{n} {w}"


def recommendations(d, st):
    """Подсказки «что сделать», собранные из показателей. Каждая — с действием, выполняемым одним нажатием."""
    out = []
    k, g = d["kpi"], d["goals"]
    wl = {w.lower() for w in (st.get("whitelist") or "").split()}
    # 1. правила, которые чаще ошибаются, чем помогают
    for r in d["rules"]:
        dec = r["accepted"] + r["rejected"]
        if dec >= 10 and r["precision"] is not None and r["precision"] < 50 and r["setting"]:
            sch = core.SCHEMA.get(r["setting"], {})
            act = {"type": "setting", "key": r["setting"], "value": False, "label": "Выключить правило"} \
                if sch.get("type") == "bool" and st.get(r["setting"]) else {"type": "link", "href": f"/settings#set-{r['setting']}", "label": "Открыть настройку"}
            out.append({"level": "warning", "title": f"Правило «{r['label']}» чаще ошибается, чем помогает",
                        "text": f"Точность {r['precision']} %: принято {r['accepted']}, отклонено {r['rejected']}. "
                                "Смягчите или отключите правило.", "action": act})
    # 2. слова, которые корректор отклоняет — кандидаты в белый список
    words = [u["word"] for u in d["unknown"] if u["rejected"] >= 2 and not u["accepted"] and not u["whitelisted"]
             and u["word"].lower() not in wl][:20]
    if words:
        out.append({"level": "info", "title": f"{plural(len(words), 'слово', 'слова', 'слов')} корректор отклоняет как «вне словаря» — "
                    "вероятно, это не ошибки",
                    "text": ", ".join(words), "action": {"type": "whitelist", "words": words, "label": "Добавить в белый список"}})
    # 3. повторяющиеся ошибки в терминах — в глоссарий
    gl = (st.get("glossary") or "").lower()
    term = re.compile(r"^[^\W\d_][\w\- ]*$")
    pairs = [r for r in d["recurring"] if r["docs"] >= 3 and r["good"] and r["rule"] in ("ai", "lt")
             and r["type"] in ("Орфография", "Единообразие", "Грамматика") and r["bad"].lower() not in gl
             and term.match(r["good"]) and term.match(r["bad"]) and r["good"].lower() != r["bad"].lower()
             and len(r["good"].split()) == len(r["bad"].split())][:5]
    for r in pairs:
        out.append({"level": "info", "title": f"«{r['bad']}» → «{r['good']}» встречается в {r['docs']} "
                                              f"{'документе' if r['docs'] % 10 == 1 and r['docs'] % 100 != 11 else 'документах'}",
                    "text": "Добавьте пару в глоссарий: ошибка будет находиться всегда, даже без языковой модели.",
                    "action": {"type": "glossary", "good": r["good"], "bad": r["bad"], "label": "В глоссарий"}})
    # 4. цели
    if goal_state("density", k["density"], g) == "bad":
        out.append({"level": "critical", "title": f"Плотность ошибок {k['density']} выше цели {g['density']}",
                    "text": f"Документов с превышением: {k.get('over_goal', 0)}. Самые проблемные — в графике «Плотность ошибок по документам».",
                    "action": {"type": "link", "href": "#c-density", "label": "Показать документы"}})
    if goal_state("author_share", k["author_share"], g) == "bad":
        out.append({"level": "warning", "title": f"Доля «Требует проверки автором» {k['author_share']} % выше цели {g['author_share']} %",
                    "text": "Пополните белый список и глоссарий; при возможности подключите языковую модель.",
                    "action": {"type": "link", "href": "#c-unknown", "label": "Слова вне словаря"}})
    if goal_state("precision", k["precision"], g) == "bad" and (k["accepted"] + k["rejected"]) >= 20:
        out.append({"level": "warning", "title": f"Точность замечаний {k['precision']} % ниже цели {g['precision']} %",
                    "text": "Посмотрите «Точность правил»: правила с пометкой ⚠ дают больше всего ложных срабатываний.",
                    "action": {"type": "link", "href": "#c-ruletab", "label": "Точность правил"}})
    if goal_state("verify_rate", k["verify_rate"], g) == "bad":
        out.append({"level": "critical", "title": f"Только {k['verify_rate']} % отчетов прошли 25 проверок",
                    "text": "Непринятые отчеты не выдаются. Откройте документы с непройденной проверкой и поправьте замечания.",
                    "action": {"type": "link", "href": "#c-verify", "label": "Машинная проверка"}})
    # 5. прочее
    if k["findings"] >= 50 and (k["reviewed_share"] or 0) < 20:
        out.append({"level": "info", "title": f"Корректор отметил только {k['reviewed_share'] or 0} % замечаний",
                    "text": "Без отметок ✓ и ✕ точность правил считается по малой выборке. Отмечайте замечания массово в обозревателе.",
                    "action": {"type": "link", "href": "#explorer", "label": "Обозреватель"}})
    if (k["ai_share"] or 0) == 0 and (k["author_share"] or 0) > 30 and st.get("ai_provider") in (None, "", "none"):
        out.append({"level": "info", "title": "Языковая модель не подключена",
                    "text": "Без нее не находятся смысловые, грамматические и логические ошибки; много мест уходит автору.",
                    "action": {"type": "link", "href": "/settings#ai", "label": "Подключить"}})
    if k.get("failed"):
        out.append({"level": "warning", "title": f"Заданий с ошибкой: {k['failed']}",
                    "text": "Проверка не выполнена — откройте задание и посмотрите журнал.",
                    "action": {"type": "link", "href": "/jobs", "label": "Задания"}})
    rank = {"critical": 0, "warning": 1, "info": 2}
    for r in out:                                      # десятичная запятая в тексте
        r["title"] = re.sub(r"(\d)\.(\d)", r"\1,\2", r["title"])
        r["text"] = re.sub(r"(\d)\.(\d)", r"\1,\2", r["text"])
    return sorted(out, key=lambda r: rank[r["level"]])


def type_goals(st=None):
    st = st or core.load_settings()
    out = {}
    for ln in (st.get("goal_density_by_type") or "").splitlines():
        if ":" in ln:
            k, v = ln.rsplit(":", 1)
            try:
                out[k.strip().lower()] = float(v.strip().replace(",", "."))
            except ValueError:
                pass
    return out


def density_goal(doc_type, st=None):
    """Цель по плотности ошибок для документа: по его типу или общая."""
    st = st or core.load_settings()
    return type_goals(st).get((doc_type or "").strip().lower()) or float(st.get("goal_density") or 0) or None


def goals(st=None):
    st = st or core.load_settings()
    return {k: (float(st.get("goal_" + k) or 0) or None) for k in ("density", "author_share", "precision", "verify_rate")}


def goal_state(key, value, g):
    """'ok' | 'bad' | None — выполнена ли цель (плотность и доля «автору» — не больше, остальное — не меньше)."""
    if value is None or not g.get(key):
        return None
    return "ok" if (value <= g[key] if key in ("density", "author_share") else value >= g[key]) else "bad"


def group_by(docs, f, field):
    """Показатели по авторам или отделам."""
    groups = collections.defaultdict(list)
    for s in docs:
        groups[s.get(field) or ""].append(s)
    out = []
    for name, ss in groups.items():
        m = metrics(ss, f)
        out.append({"key": name, "label": name or "не указан", "docs": m["docs"], "pages": m["pages"],
                    "findings": m["findings"], "density": m["density"], "precision": m["precision"],
                    "author_share": m["author_share"], "verify_rate": m["verify_rate"]})
    out.sort(key=lambda r: (r["key"] == "", -(r["docs"])))
    return out


def options(stats):
    fam = collections.Counter(s["family"] for s in stats)
    batches = collections.OrderedDict()
    for s in stats:
        if s["batch"]:
            batches.setdefault(s["batch"], {"id": s["batch"], "n": 0, "first": s["name"], "date": s["created"][:10]})
            batches[s["batch"]]["n"] += 1
    return {
        "types": TYPES,
        "sources": [{"key": k, "label": v} for k, v in SOURCES.items()],
        "statuses": [{"key": k, "label": v} for k, v in STATUS.items()],
        "families": [{"key": k, "n": n} for k, n in sorted(fam.items())],
        "batches": [b for b in reversed(list(batches.values())) if b["n"] > 1][:60],
        "authors": sorted({s.get("author") for s in stats if s.get("author")}),
        "departments": sorted({s.get("department") for s in stats if s.get("department")}),
        "doc_types": sorted({s.get("doc_type") for s in stats if s.get("doc_type")}),
        "first": stats[0]["created"][:10] if stats else None,
        "last": stats[-1]["created"][:10] if stats else None,
    }


# ------------------------------------------------------------------ обозреватель замечаний
@_limited
def explore(q, limit=200, offset=0):
    f = parse_filters(q)
    text = (q.get("text") or "").strip().lower()
    word = (q.get("bad") or "").strip().lower()
    docs = filter_docs(all_stats(), f)
    rows, total = [], 0
    for s in reversed(docs):
        for x in s["findings"]:
            if not finding_ok(x, f):
                continue
            if word and x["bad"].lower() != word:
                continue
            if text and text not in (x["bad"] + " " + x["good"]).lower():
                continue
            if offset <= total < offset + limit:     # строки собираются только для запрошенной страницы
                rows.append(dict(x, job=s["id"], name=s["name"], created=s["created"][:10],
                                 rule_label=rule_info(x["rule"])["label"], source_label=SOURCES.get(x["source"], "")))
            total += 1
    return {"total": total, "rows": rows, "offset": offset, "limit": limit}


# ------------------------------------------------------------------ тренды
def bucket_of(day, gran):
    if gran == "day":
        return day
    if gran == "week":
        return day - dt.timedelta(days=day.weekday())
    return day.replace(day=1)


def bucket_next(b, gran):
    if gran == "day":
        return b + dt.timedelta(days=1)
    if gran == "week":
        return b + dt.timedelta(days=7)
    return (b.replace(day=28) + dt.timedelta(days=4)).replace(day=1)


MONTHS = ["янв", "фев", "мар", "апр", "май", "июн", "июл", "авг", "сен", "окт", "ноя", "дек"]


def bucket_label(b, gran):
    if gran == "month":
        return f"{MONTHS[b.month - 1]} {b.year}"
    if gran == "week":
        return f"нед. {b:%d.%m}"
    return f"{b:%d.%m}"


def linfit(ys):
    """Линейный тренд по точкам (индекс, значение); пропуски игнорируются. → (a, b) для y = a + b·x."""
    pts = [(i, y) for i, y in enumerate(ys) if y is not None]
    if len(pts) < 2:
        return None
    n = len(pts)
    mx = sum(p[0] for p in pts) / n
    my = sum(p[1] for p in pts) / n
    sxx = sum((p[0] - mx) ** 2 for p in pts)
    if not sxx:
        return None
    b = sum((p[0] - mx) * (p[1] - my) for p in pts) / sxx
    return my - b * mx, b


def moving_avg(ys, w=3):
    out = []
    for i in range(len(ys)):
        win = [y for y in ys[max(0, i - w + 1):i + 1] if y is not None]
        out.append(_r(sum(win) / len(win)) if win else None)
    return out


BETTER = {"docs": 0, "pages": 0, "findings": 0, "density": -1, "author_share": -1, "precision": 1,
          "verify_rate": 1, "avg_duration": -1, "reviewed_share": 1, "per_doc": -1, "sec_per_page": -1}


def compare(cur, prev):
    out = {}
    for k, better in BETTER.items():
        a, b = cur.get(k), prev.get(k)
        d = None
        if a is not None and b is not None:
            d = _r(a - b, 2)
        pct = _r(100.0 * (a - b) / b, 1) if (a is not None and b) else None
        verdict = "same"
        if d is not None and better and abs(d) > 1e-9:
            verdict = "better" if (d > 0) == (better > 0) else "worse"
        elif d is not None and abs(d) > 1e-9:
            verdict = "up" if d > 0 else "down"
        out[k] = {"cur": a, "prev": b, "delta": d, "pct": pct, "verdict": verdict}
    return out


@_limited
def trends(q):
    f = parse_filters(dict(q, period=q.get("period") or "180"))
    gran = q.get("gran") if q.get("gran") in ("day", "week", "month") else "week"
    stats = all_stats()
    pool = filter_docs(stats, dict(f, **{"from": None, "to": None}))      # без ограничения дат: для «нового» и сравнения
    if f["from"] is None:
        f["from"] = _d(pool[0]["created"]) if pool else dt.date.today()
    if f["to"] is None:
        f["to"] = dt.date.today()
    docs = [s for s in pool if _in_range(s["created"], f["from"], f["to"])]

    # корзины
    buckets = []
    b = bucket_of(f["from"], gran)
    while b <= f["to"]:
        buckets.append(b)
        b = bucket_next(b, gran)
    buckets = buckets[-800:]                    # длинный период по дням: отбрасываются самые старые, а не свежие
    idx = {bk: i for i, bk in enumerate(buckets)}
    groups = [[] for _ in buckets]
    for s in docs:
        i = idx.get(bucket_of(_d(s["created"]), gran))
        if i is not None:
            groups[i].append(s)
    per = [metrics(g, f) for g in groups]

    def ser(k):
        return [p.get(k) for p in per]

    dens = ser("density")
    fit = linfit(dens)
    horizon = 3
    first = next((i for i, y in enumerate(dens) if y is not None), 0)
    trend = [_r(max(0.0, fit[0] + fit[1] * i)) if fit and i >= first else None for i in range(len(buckets))]
    fb = buckets[-1] if buckets else None
    future = []
    for _ in range(horizon):
        if fb is None:
            break
        fb = bucket_next(fb, gran)
        future.append(fb)
    forecast = [_r(max(0.0, fit[0] + fit[1] * (len(buckets) + i))) if fit else None for i in range(len(future))]

    by_type = {t: [sum(1 for s in g for x in s["findings"] if x["type"] == t and finding_ok(x, f)) for g in groups]
               for t in TYPES}
    by_source = {}
    for k in SOURCES:
        vals = [sum(1 for s in g for x in s["findings"] if x["source"] == k and finding_ok(x, f)) for g in groups]
        if any(vals):
            by_source[k] = vals

    # сравнение с предыдущим равным периодом
    span = (f["to"] - f["from"]).days + 1
    pfrom, pto = f["from"] - dt.timedelta(days=span), f["from"] - dt.timedelta(days=1)
    prev_docs = [s for s in pool if _in_range(s["created"], pfrom, pto)]
    cmp_ = compare(metrics(docs, f), metrics(prev_docs, f))

    # тренды правил
    rule_tot = collections.Counter(x["rule"] for s in docs for x in s["findings"] if finding_ok(x, f))
    prev_rule = collections.Counter(x["rule"] for s in prev_docs for x in s["findings"] if finding_ok(x, f))
    rule_rows = []
    for rule in sorted(set(rule_tot) | set(prev_rule), key=lambda r: -(rule_tot.get(r, 0) + prev_rule.get(r, 0) / 2)):
        ri = rule_info(rule)
        spark = [sum(1 for s in g for x in s["findings"] if x["rule"] == rule and finding_ok(x, f)) for g in groups]
        dens_spark = [_r(10.0 * spark[i] / per[i]["pages"]) if per[i]["pages"] else None for i in range(len(groups))]
        cur, prv = rule_tot.get(rule, 0), prev_rule.get(rule, 0)
        cur_d = 10.0 * cur / max(1, sum(s["pages"] for s in docs))
        prv_d = 10.0 * prv / max(1, sum(s["pages"] for s in prev_docs)) if prev_docs else None
        lf = linfit(dens_spark)
        rule_rows.append(dict(ri, cur=cur, prev=prv, cur_density=_r(cur_d, 3), prev_density=_r(prv_d, 3),
                              delta_pct=_r(100 * (cur_d - prv_d) / prv_d, 1) if prv_d else None,
                              slope=_r(lf[1], 4) if lf else None, spark=spark))

    # новое в периоде: правила, слова и пары ошибок, которых раньше не было
    before = [s for s in pool if _d(s["created"]) < f["from"]]
    seen_rules = {x["rule"] for s in before for x in s["findings"]}
    seen_words = {x["bad"].lower() for s in before for x in s["findings"] if x["rule"] == "spell"}
    seen_pairs = {(x["bad"].lower(), x["good"].lower()) for s in before for x in s["findings"]}
    new_rules = collections.Counter(x["rule"] for s in docs for x in s["findings"]
                                    if x["rule"] not in seen_rules and finding_ok(x, f))
    new_words = collections.Counter(x["bad"] for s in docs for x in s["findings"]
                                    if x["rule"] == "spell" and x["bad"].lower() not in seen_words)
    new_pairs = collections.Counter((x["bad"], x["good"], x["type"]) for s in docs for x in s["findings"]
                                    if x["rule"] not in ("spell", "batch") and x["bad"]
                                    and (x["bad"].lower(), x["good"].lower()) not in seen_pairs and finding_ok(x, f))
    gone = sorted(r for r in seen_rules - {x["rule"] for s in docs for x in s["findings"]} if r)

    # семейства документов (несколько версий одного документа)
    fams = collections.defaultdict(list)
    for s in pool:
        fams[s["family"]].append(s)
    families = []
    for fam, ss in fams.items():
        if len(ss) < 2:
            continue
        ss.sort(key=lambda s: (s["created"], s["id"]))
        vers = [{"id": s["id"], "name": s["name"], "created": s["created"][:10], "findings": len(s["findings"]),
                 "density": _r(10.0 * len(s["findings"]) / s["pages"]), "verify_ok": s["verify_ok"]} for s in ss]
        a, z = vers[0]["density"], vers[-1]["density"]
        families.append({"family": fam, "versions": vers, "first": a, "last": z,
                         "change_pct": _r(100 * (z - a) / a, 1) if a else None})
    families.sort(key=lambda r: -len(r["versions"]))

    # календарь активности: 53 недели до конца периода
    cal_end = f["to"]
    cal_start = bucket_of(cal_end - dt.timedelta(days=364), "week")
    daily = collections.Counter(_d(s["created"]) for s in pool)
    daily_f = collections.Counter()
    for s in pool:
        daily_f[_d(s["created"])] += sum(1 for x in s["findings"] if finding_ok(x, f))
    cal = []
    day = cal_start
    while day <= cal_end:
        cal.append({"d": day.isoformat(), "docs": daily.get(day, 0), "findings": daily_f.get(day, 0)})
        day += dt.timedelta(days=1)

    # нагрузка: день недели × час
    wh = [[0] * 24 for _ in range(7)]
    for s in docs:
        try:
            t = dt.datetime.fromisoformat(s["created"])
            wh[t.weekday()][t.hour] += 1
        except ValueError:
            pass

    return {
        "filters": {k: (v.isoformat() if isinstance(v, dt.date) else v) for k, v in f.items() if k != "owner"},
        "gran": gran,
        "options": options([s for s in stats if f.get("owner") is None or _visible(f["owner"], s)]),
        "labels": [bucket_label(b, gran) for b in buckets],
        "keys": [b.isoformat() for b in buckets],
        "future_labels": [bucket_label(b, gran) for b in future],
        "series": {k: ser(k) for k in ("docs", "pages", "findings", "occurrences", "density", "author_share",
                                       "precision", "verify_rate", "avg_duration", "sec_per_page", "ai_share",
                                       "reviewed_share", "per_doc")},
        "density_ma": moving_avg(dens),
        "density_trend": trend,
        "density_forecast": forecast,
        "density_slope": _r(fit[1], 4) if fit else None,
        "by_type": by_type,
        "by_source": {SOURCES[k]: v for k, v in by_source.items()},
        "compare": cmp_,
        "prev_period": {"from": pfrom.isoformat(), "to": pto.isoformat(), "docs": len(prev_docs)},
        "rules": rule_rows[:25],
        "new_rules": [dict(rule_info(r), n=n) for r, n in new_rules.most_common()],
        "new_words": [{"word": w, "n": n} for w, n in new_words.most_common(30)],
        "new_pairs": [{"bad": b, "good": g, "type": t, "n": n} for (b, g, t), n in new_pairs.most_common(30)],
        "gone_rules": [rule_info(r) for r in gone],
        "families": families[:20],
        "calendar": cal,
        "weekday_hour": wh,
        "has_before": bool(before),
        "goals": goals(),
    }


# ------------------------------------------------------------------ экспорт
def export_rows(data):
    """Таблицы для выгрузки XLSX/CSV: [(название листа, заголовки, строки)]."""
    k = data["kpi"]
    sheets = [
        ("Показатели", ["Показатель", "Значение"], [
            ["Документов", k["docs"]], ["Страниц", k["pages"]], ["Замечаний (строк отчета)", k["findings"]],
            ["Мест с ошибками (с учетом системных)", k["occurrences"]], ["Замечаний на 10 страниц", k["density"]],
            ["Замечаний на документ", k["per_doc"]], ["Доля «Требует проверки автором», %", k["author_share"]],
            ["Доля замечаний ИИ, %", k["ai_share"]], ["Принято", k["accepted"]], ["Отклонено", k["rejected"]],
            ["Не просмотрено", k["new"]], ["Точность (принято / просмотрено), %", k["precision"]],
            ["Просмотрено корректором, %", k["reviewed_share"]], ["Отредактировано", k["edited"]],
            ["Отчетов, прошедших 25 проверок, %", k["verify_rate"]], ["Среднее время обработки, с", k["avg_duration"]],
            ["Медиана времени, с", k["median_duration"]], ["Секунд на страницу", k["sec_per_page"]],
            ["Кандидатов, отсеянных машинно", k["auto_rejected"]], ["Заданий с ошибкой", k["failed"]]]),
        ("По типам", ["Тип", "Замечаний"], [[r["label"], r["n"]] for r in data["by_type"]]),
        ("По источникам", ["Источник", "Замечаний"], [[r["label"], r["n"]] for r in data["by_source"]]),
        ("Правила", ["Правило", "Код", "Источник", "Замечаний", "Документов", "Принято", "Отклонено",
                     "Не просмотрено", "Точность, %", "Настройка"],
         [[r["label"], r["rule"], r["source_label"], r["n"], r["docs"], r["accepted"], r["rejected"], r["new"],
           r["precision"], r["setting"]] for r in data["rules"]]),
        ("Документы", ["№", "Документ", "Дата", "Страниц", "Замечаний", "На 10 страниц", "25/25"],
         [[r["id"], r["name"], r["created"], r["pages"], r["findings"], r["density"], "да" if r["verify_ok"] else "нет"]
          for r in data["density"]]),
        ("Документ × тип", ["№", "Документ"] + data["heatmap"]["cols"],
         [[r["id"], r["name"]] + r["values"] for r in data["heatmap"]["rows"]]),
        ("Положение", ["Часть документа", "Замечаний"], [[r["label"], r["n"]] for r in data["position"]]),
        ("Повторяющиеся", ["Ошибка", "Как правильно", "Тип", "Правило", "Мест", "Документов"],
         [[r["bad"], r["good"], r["type"], r["rule_label"], r["n"], r["docs"]] for r in data["recurring"]]),
        ("Слова вне словаря", ["Слово", "Раз", "Документов", "Принято", "Отклонено", "В белом списке"],
         [[r["word"], r["n"], r["docs"], r["accepted"], r["rejected"], "да" if r["whitelisted"] else "нет"]
          for r in data["unknown"]]),
        ("Предпечатная", ["Проблема", "Документов", "Случаев"], [[r["label"], r["docs"], r["items"]] for r in data["preflight"]]),
        ("Машинная проверка", ["Не пройдено требование", "Документов"], [[r["label"], r["n"]] for r in data["verify_fail"]]),
        ("Время обработки", ["№", "Документ", "Страниц", "Секунд"], [[r["id"], r["name"], r["x"], r["y"]] for r in data["duration"]]),
        ("По авторам", ["Автор", "Документов", "Страниц", "Замечаний", "На 10 страниц", "Точность, %", "«Автору», %", "25/25, %"],
         [[r["label"], r["docs"], r["pages"], r["findings"], r["density"], r["precision"], r["author_share"], r["verify_rate"]]
          for r in data["by_author"]]),
        ("По отделам", ["Отдел", "Документов", "Страниц", "Замечаний", "На 10 страниц", "Точность, %", "«Автору», %", "25/25, %"],
         [[r["label"], r["docs"], r["pages"], r["findings"], r["density"], r["precision"], r["author_share"], r["verify_rate"]]
          for r in data["by_department"]]),
        ("Рекомендации", ["Важность", "Что", "Подробности"],
         [[{"critical": "высокая", "warning": "средняя", "info": "низкая"}[r["level"]], r["title"], r["text"]]
          for r in data["recommendations"]]),
    ]
    return sheets


def trend_rows(t):
    head = ["Период"] + ["Документов", "Страниц", "Замечаний", "На 10 страниц", "Скользящее среднее", "Линейный тренд",
                         "Доля автору, %", "Точность, %", "25/25, %", "Среднее время, с"]
    s = t["series"]
    rows = [[t["labels"][i], s["docs"][i], s["pages"][i], s["findings"][i], s["density"][i], t["density_ma"][i],
             t["density_trend"][i], s["author_share"][i], s["precision"][i], s["verify_rate"][i], s["avg_duration"][i]]
            for i in range(len(t["labels"]))]
    rows += [[lab + " (прогноз)", None, None, None, None, None, t["density_forecast"][i], None, None, None, None]
             for i, lab in enumerate(t["future_labels"])]
    types = [["Период"] + list(t["by_type"])] + [[t["labels"][i]] + [t["by_type"][k][i] for k in t["by_type"]]
                                                 for i in range(len(t["labels"]))]
    rules = [[r["label"], r["rule"], r["cur"], r["prev"], r["cur_density"], r["prev_density"], r["delta_pct"], r["slope"]]
             for r in t["rules"]]
    fams = [[f["family"], v["name"], v["created"], v["findings"], v["density"]] for f in t["families"] for v in f["versions"]]
    return [
        ("Динамика", head, rows),
        ("Типы по периодам", types[0], types[1:]),
        ("Правила", ["Правило", "Код", "Замечаний", "В пред. периоде", "На 10 стр.", "На 10 стр. ранее",
                     "Изменение, %", "Наклон тренда"], rules),
        ("Версии документов", ["Семейство", "Документ", "Дата", "Замечаний", "На 10 страниц"], fams),
        ("Новое в периоде", ["Ошибка", "Как правильно", "Тип", "Раз"],
         [[r["bad"], r["good"], r["type"], r["n"]] for r in t["new_pairs"]]),
    ]


def to_xlsx(sheets, title):
    """Книга Excel в режиме потоковой записи (write_only): 100 тысяч строк не держатся в памяти целиком."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook(write_only=True)
    font, fill, al = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="0F4382"), Alignment(vertical="center", wrap_text=True)
    for name, head, rows in sheets:
        rows = rows if isinstance(rows, list) else list(rows)
        ws = wb.create_sheet(re.sub(r"[\[\]:*?/\\]", "", name)[:31])
        for i, h in enumerate(head, 1):              # ширина колонок — до записи строк
            width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:300] if i - 1 < len(r) and r[i - 1] is not None])
            ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, width + 2))
        ws.freeze_panes = "A2"
        cells = []
        for h in head:
            c = WriteOnlyCell(ws, value=h)
            c.font, c.fill, c.alignment = font, fill, al
            cells.append(c)
        ws.append(cells)
        for r in rows:
            ws.append([("" if v is None else safe(v)) for v in r])
    wb.properties.title = title
    import io
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_csv(sheets):
    import csv
    import io
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    for name, head, rows in sheets:
        w.writerow([f"# {name}"])
        w.writerow(head)
        for r in rows:
            w.writerow(["" if v is None else (str(v).replace(".", ",") if isinstance(v, float) else safe(v)) for v in r])
        w.writerow([])
    return "﻿" + buf.getvalue()


def nan_safe(o):
    """JSON без NaN/inf."""
    if isinstance(o, float) and (math.isnan(o) or math.isinf(o)):
        return None
    if isinstance(o, dict):
        return {k: nan_safe(v) for k, v in o.items()}
    if isinstance(o, list):
        return [nan_safe(v) for v in o]
    return o
