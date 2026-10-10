# -*- coding: utf-8 -*-
"""Отчет прямо в исходном документе Word: примечания и исправления в режиме рецензирования.

Для документа, загруженного в формате DOCX, создается копия, в которой:
  * каждое замечание — примечание Word (тип, как правильно, пояснение) на ошибочном фрагменте;
  * ошибки с готовым исправлением — правка в режиме рецензирования (удаление «как было» + вставка «как правильно»),
    автор принимает ее кнопкой «Принять» на вкладке «Рецензирование».
«Требует проверки автором» — только примечание, без правки.
"""
import copy
import datetime as dt
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from .common import AUTHOR

W_T = qn("w:t")


def _paragraphs(doc):
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        yield from _table_pars(t)


def _table_pars(t):
    for row in t.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nt in cell.tables:
                yield from _table_pars(nt)


def _runs(p):
    """Прямые фрагменты текста абзаца (без уже вставленных правок)."""
    return [r for r in p._p.iter(qn("w:r")) if r.getparent() is p._p or r.getparent().tag == qn("w:hyperlink")]


SIMPLE_CHILDREN = {qn("w:rPr"), W_T}


def _simple(r):
    """Фрагмент из одного текста (w:t) и оформления: такой можно безопасно делить и править.
    Табуляции, переносы, рисунки, надписи, сноски и поля не трогаем."""
    kids = list(r)
    return bool(kids) and all(k.tag in SIMPLE_CHILDREN for k in kids) and len(r.findall(W_T)) == 1


def _text(r):
    """Видимый текст фрагмента; табуляции и переносы — как пробел (для поиска), содержимое надписей не берется."""
    out = []
    for k in r:
        if k.tag == W_T:
            out.append(k.text or "")
        elif k.tag in (qn("w:tab"), qn("w:br"), qn("w:cr")):
            out.append(" ")
        elif k.tag == qn("w:noBreakHyphen"):
            out.append("-")
        elif k.tag in (qn("w:drawing"), qn("w:pict"), qn("w:object"), qn("w:footnoteReference"),
                       qn("w:endnoteReference")) or k.tag.endswith("}AlternateContent"):
            out.append("\uFFFC")                      # непрозрачный объект — разрыв для поиска
    return "".join(out)


def _split(r, offset):
    """Делит простой фрагмент r на два по позиции offset; возвращает правую часть."""
    txt = r.find(W_T).text or ""
    right = copy.deepcopy(r)
    lt, rt = r.find(W_T), right.find(W_T)
    lt.text, rt.text = txt[:offset], txt[offset:]
    for t in (lt, rt):
        t.set(qn("xml:space"), "preserve")
    r.addnext(right)
    return right


def _isolate(p, start, end):
    """Фрагменты абзаца, точно покрывающие символы [start, end). None — если текст разбит неудобно (поля и т. п.)."""
    pos = 0
    runs = _runs(p)
    out = []
    for r in runs:
        if r.find(qn("w:fldChar")) is not None or r.find(qn("w:instrText")) is not None:
            continue
        txt = _text(r)
        n = len(txt)
        if not n:
            continue
        a, b = pos, pos + n
        pos = b
        if b <= start or a >= end:
            continue
        if not _simple(r):
            return None                                # ошибка задевает табуляцию, рисунок и т. п. — не правим
        if a < start:
            r = _split(r, start - a)
            a = start
        if b > end:
            _split(r, end - a)
        out.append(r)
    return out or None


def _ptext(p):
    return "".join(_text(r) for r in _runs(p) if r.find(qn("w:fldChar")) is None)


def _norm_ws(s):
    return re.sub(r"\s+", " ", (s or "").replace(" ", " ")).strip()


class _Texts:
    """Тексты абзацев: извлекаются один раз, после правки обновляется только измененный абзац (прежде текст всех
    абзацев документа извлекался заново для каждого замечания — минуты и часы на больших документах)."""

    def __init__(self, pars):
        self.pars = pars
        self.raw = [_ptext(p) for p in pars]
        self.soft = [t.replace("\xa0", " ") for t in self.raw]
        self.pos = {}                                  # объединенная ячейка таблицы — один абзац в нескольких местах
        for i, p in enumerate(pars):
            self.pos.setdefault(id(p._p), []).append(i)

    def refresh(self, p):
        text = _ptext(p)
        for i in self.pos.get(id(p._p), []):
            self.raw[i] = text
            self.soft[i] = text.replace("\xa0", " ")


def _find(pars, f, texts=None):
    """Абзац и позиция ошибки: сначала по цитате (контекст), затем по самой ошибке."""
    bad = f.get("bad") or ""
    if not bad.strip():
        return None
    src = f.get("src") or bad
    texts = texts or _Texts(pars)
    for exact in (True, False):
        for k, p in enumerate(pars):
            t = texts.raw[k]
            if not t:
                continue
            hay, needle_src, needle = (t, src, bad) if exact else (texts.soft[k], src.replace("\xa0", " "), bad.replace("\xa0", " "))
            i = hay.find(needle_src)
            if i >= 0:
                j = hay.find(needle, i, i + len(needle_src) + len(needle))
                if j >= 0:
                    return p, j, j + len(bad)
            if exact:
                continue
            j = hay.find(needle)
            if j >= 0 and len(needle.strip()) >= 3:
                return p, j, j + len(bad)
    return None


def _rev(tag, rid, author, date):
    e = OxmlElement(tag)
    e.set(qn("w:id"), str(rid))
    e.set(qn("w:author"), author)
    e.set(qn("w:date"), date)
    return e


def review(src_path, findings, out_path, author="Корректор"):
    """→ (примечаний, исправлений, не найдено)."""
    doc = Document(src_path)
    pars = list(_paragraphs(doc))
    date = dt.datetime.now().replace(microsecond=0).isoformat() + "Z"
    rid = 9000
    n_comments = n_changes = missing = 0
    items = []
    for i, f in enumerate(sorted(findings, key=lambda f: f.get("page", 0)), 1):
        bads = [b for b in (f.get("bad") or "").split(" … ") if b.strip()]
        if len(bads) > 1:                              # объединенное замечание: несколько мест в одной цитате
            goods = (f.get("good") or "").split(" … ")
            same = len(goods) == len(bads)
            for k, b in enumerate(bads):
                items.append((i, dict(f, bad=b, good=goods[k] if same else "—", _comment_only=not same)))
        else:
            items.append((i, f))
    texts = _Texts(pars)
    for i, f in items:
        loc = _find(pars, f, texts)
        if not loc:
            missing += 1
            continue
        p, s, e = loc
        runs = _isolate(p, s, e)
        if not runs:
            missing += 1
            continue
        good = f.get("good") or ""
        text = f"{(f.get('type') or '').rstrip('.')}. "
        if f.get("type") != AUTHOR and f.get("fix") not in (None, "", "—"):
            text += f"Как правильно: {f['fix'].rstrip(' .')}. "
        text += f.get("note") or ""
        first, last = runs[0], runs[-1]
        if f.get("type") != AUTHOR and not f.get("_comment_only") and good not in ("", "—") and good != f.get("bad"):
            # правка: <w:del>как было</w:del><w:ins>как правильно</w:ins>
            rid += 1
            d = _rev("w:del", rid, author, date)
            first.addprevious(d)
            for r in runs:
                d.append(r)
                for t in r.findall(W_T):
                    t.tag = qn("w:delText")
            rid += 1
            ins = _rev("w:ins", rid, author, date)
            nr = OxmlElement("w:r")                    # новый фрагмент: оформление первого + текст исправления
            rpr = runs[0].find(qn("w:rPr"))
            if rpr is not None:
                nr.append(copy.deepcopy(rpr))
            t0 = OxmlElement("w:t")
            t0.text = good
            t0.set(qn("xml:space"), "preserve")
            nr.append(t0)
            ins.append(nr)
            d.addnext(ins)
            first_el, last_el = d, ins
            n_changes += 1
        else:
            first_el, last_el = first, last
        c = doc.comments.add_comment(text=f"№ {i}. {text}".strip(), author=author, initials="К")
        cs = OxmlElement("w:commentRangeStart")
        cs.set(qn("w:id"), str(c.comment_id))
        first_el.addprevious(cs)
        ce = OxmlElement("w:commentRangeEnd")
        ce.set(qn("w:id"), str(c.comment_id))
        last_el.addnext(ce)
        ref = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        st = OxmlElement("w:rStyle")
        st.set(qn("w:val"), "CommentReference")
        rpr.append(st)
        ref.append(rpr)
        cr = OxmlElement("w:commentReference")
        cr.set(qn("w:id"), str(c.comment_id))
        ref.append(cr)
        ce.addnext(ref)
        texts.refresh(p)                               # удаленный текст правки больше не ищется в абзаце
        n_comments += 1
    doc.save(out_path)
    return n_comments, n_changes, missing
