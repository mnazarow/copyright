# build.py v2 — собирает отчет DOCX из findings.json строго по требованиям оформления.
# Запуск:  python build.py findings.json out.docx
import sys, json
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

IN = sys.argv[1] if len(sys.argv) > 1 else "findings.json"
OUT = sys.argv[2] if len(sys.argv) > 2 else "otchet.docx"
NAVY, RED, GREEN, GREY = "0F4382", RGBColor(0xC0, 0, 0), RGBColor(0x00, 0x80, 0x00), "B2B2B2"
# маркер = часть цитаты, которую надо выделить. Несколько частей разделяются « … ».
# Прочерк «—» означает «выделять нечего» (ячейка остается черной).

data = json.load(open(IN, encoding="utf-8"))
meta, F = data["document"], data["findings"]
F = sorted(F, key=lambda r: (int(r["page"]) if str(r["page"]).isdigit() else 999, str(r["clause"])))

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Arial"; st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
st.paragraph_format.space_after = Pt(4)
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, m, Cm(1.2))


def heading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = "Arial"; r.font.size = Pt(14); r.bold = True
    r.font.color.rgb = RGBColor(0x0F, 0x43, 0x82)


def borders_only_inside(table):
    b = OxmlElement("w:tblBorders")
    for e in ("top", "left", "bottom", "right"):
        x = OxmlElement("w:" + e); x.set(qn("w:val"), "nil"); b.append(x)
    for e in ("insideH", "insideV"):
        x = OxmlElement("w:" + e)
        x.set(qn("w:val"), "single"); x.set(qn("w:sz"), "6")
        x.set(qn("w:space"), "0"); x.set(qn("w:color"), GREY)
        b.append(x)
    table._tbl.tblPr.append(b)


def fill(cell, color):
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(s)


def split_marker(marker):
    """«— 10 лет» → «— 10 лет»; «50мм … 100мм» → ['50мм','100мм']; «—» → []"""
    if not marker:
        return []
    m = marker.strip()
    if m in {"—", "-", "–", ""}:
        return []
    return [p.strip() for p in m.split(" … ") if p.strip() and p.strip() not in {"—", "-", "–"}]


def rich(par, text, marker, color):
    """Пишет text в ячейку, выделяя marker полужирным заданного цвета.
    Если маркер не найден — пишет целиком обычным шрифтом (это поймает verify.py)."""
    parts = split_marker(marker)
    if not parts:
        par.add_run(text); return
    spans, pos = [], 0
    for part in parts:
        i = text.find(part, pos)
        if i < 0:
            i = text.find(part)          # вторая попытка: искать с начала
            if i < 0:
                continue
        spans.append((i, i + len(part)))
        pos = i + len(part)
    if not spans:
        par.add_run(text); return
    spans.sort()
    cur = 0
    for a, b in spans:
        if a < cur:
            continue
        par.add_run(text[cur:a])
        r = par.add_run(text[a:b]); r.bold = True; r.font.color.rgb = color
        cur = b
    par.add_run(text[cur:])


heading("Отчет о корректорской вычитке")
for line in meta["header_lines"]:
    doc.add_paragraph(line).paragraph_format.space_after = Pt(2)
doc.add_paragraph(meta["totals"]).paragraph_format.space_after = Pt(6)

cols = ["№", "Стр., пункт", "Тип", "Как в документе", "Как правильно", "Пояснение"]
widths = [Cm(0.8), Cm(2.6), Cm(2.9), Cm(6.5), Cm(6.5), Cm(8.0)]
t = doc.add_table(rows=1, cols=6)
t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
for i, c in enumerate(cols):
    cell = t.rows[0].cells[i]; cell.text = ""
    r = cell.paragraphs[0].add_run(c)
    r.bold = True; r.font.name = "Arial"; r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    fill(cell, NAVY)
t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))

for i, f in enumerate(F, 1):
    row = t.add_row()
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    c = row.cells
    c[0].text = str(i)
    c[1].text = f'стр. {f["page"]}, {f["clause"]}'
    c[2].text = f["type"]
    rich(c[3].paragraphs[0], f["src"], f["bad"], RED)
    rich(c[4].paragraphs[0], f["fix"], f["good"], GREEN)
    c[5].text = f["note"]
borders_only_inside(t)
for row in t.rows:
    for i, w in enumerate(widths):
        row.cells[i].width = w

heading("Итог")
doc.add_paragraph(meta["summary"])
doc.save(OUT)
print("собрано:", OUT, "| строк в таблице:", len(F))
