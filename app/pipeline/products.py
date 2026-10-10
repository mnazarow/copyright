# -*- coding: utf-8 -*-
"""Сверка с справочником продукции (выгрузка из 1С, ERP или PIM в CSV / XLSX).

Справочник: первая колонка — артикул, далее — наименование и характеристики. Единица измерения берется из
заголовка колонки: «Мощность, кВт», «Напор (м)», «Расход [м3/ч]». Колонки без единицы не сверяются.

Проверки документа:
  * артикул похож на артикул справочника, но отличается 1–2 знаками → «возможна опечатка в артикуле»;
  * артикул с префиксом серии из справочника, которого в справочнике нет → «артикул не найден в справочнике»;
  * в строке (строке таблицы) с одним артикулом справочника значение с единицей колонки отличается от
    справочного больше допуска → «расхождение с справочником».
"""
import csv
import io
import json
import os
import re
import threading
import time
from collections import Counter, defaultdict

MAX_ROWS = 300_000                     # строк справочника: больше — не выгрузка каталога, а ошибка (или «zip-бомба»)
MAX_COLS = 200

def PATH():
    return os.path.join(os.environ.get("KORREKTOR_DATA") or os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(
        os.path.abspath(__file__)))), "data"), "products.json")


LOOK = str.maketrans("АВЕКМНОРСТХУаеоркхсу", "ABEKMHOPCTXYAEOPKXCY")
TOKEN = re.compile(r"(?<![\w\-/.])(?=[\w\-/.]*\d)(?=[\w\-/.]*[A-Za-zА-Яа-яЁё])[A-Za-zА-Яа-яЁё0-9][\w\-/.]{1,38}[A-Za-zА-Яа-яЁё0-9](?![\w\-/])")
UNIT_IN_HEADER = re.compile(r"^(.*?)\s*(?:,|\(|\[)\s*([^()\[\],]{1,15}?)\s*[)\]]?\s*$")


def norm(a):
    return re.sub(r"\s+", "", str(a or "")).upper().translate(LOOK).replace("Ё", "Е")


def norm_unit(u):
    return (u or "").replace("³", "3").replace("²", "2").replace(" ", "").replace(" ", "").lower()


def num(v):
    s = str(v or "").strip().replace(" ", "").replace(" ", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def _xlsx_rows(data):
    """Строки листа по одной (прежде весь лист превращался в список строк: XLSX в 50 МБ раскрывался в гигабайты)."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        width = min(MAX_COLS, ws.max_column or MAX_COLS)   # строки дополняются до этой ширины — не шире листа
        for n, row in enumerate(ws.iter_rows(values_only=True, max_col=width)):
            if n > MAX_ROWS:
                raise ValueError(f"в справочнике больше {MAX_ROWS:,} строк".replace(",", " "))
            yield ["" if c is None else str(c) for c in row]
    finally:
        wb.close()


def _rows_from_bytes(data, filename):
    if filename.lower().endswith((".xlsx", ".xlsm")):
        return _xlsx_rows(data)
    for enc in ("utf-8-sig", "cp1251"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("файл не в UTF-8 и не в Windows-1251")
    header = text.lstrip("\ufeff").splitlines()[0] if text.strip() else ""
    # разделитель — по строке заголовков; «;» и табуляция важнее запятой (в 1С запятая — десятичный знак)
    delim = ";" if ";" in header else "\t" if "\t" in header else ","
    return list(csv.reader(io.StringIO(text), delimiter=delim))


def parse(data, filename="products.csv"):
    """→ справочник {"columns": [...], "units": {...}, "items": {нормированный артикул: {...}}, "file": имя}."""
    rows = (r for r in _rows_from_bytes(data, filename) if any(str(c).strip() for c in r))
    first = next(rows, None)
    if first is None:
        raise ValueError("в файле нет строк: нужна строка заголовков и хотя бы одна строка данных")
    head = [str(h).strip() for h in first]
    units = {}
    for i, h in enumerate(head[1:], 1):
        m = UNIT_IN_HEADER.match(h)
        if m and m.group(2).strip() and not m.group(2).strip().isdigit():
            units[i] = (m.group(1).strip() or h, m.group(2).strip())
    items = {}
    n = -1
    for n, r in enumerate(rows):
        if n > MAX_ROWS:
            raise ValueError(f"в справочнике больше {MAX_ROWS:,} строк".replace(",", " "))
        art = str(r[0]).strip() if r else ""
        if not art or len(art) > 60:
            continue
        params = {}
        for i, (name, unit) in units.items():
            if i < len(r) and num(r[i]) is not None:
                params[name] = {"value": num(r[i]), "unit": unit, "raw": str(r[i]).strip()}
        items[norm(art)] = {"article": art, "name": str(r[1]).strip() if len(r) > 1 and 1 not in units else "",
                            "params": params}
    if n < 0:
        raise ValueError("в файле нет строк: нужна строка заголовков и хотя бы одна строка данных")
    if not items:
        raise ValueError("не найдено ни одного артикула в первой колонке")
    return {"columns": head, "units": {v[0]: v[1] for v in units.values()}, "items": items,
            "file": os.path.basename(filename)}


def _write_json(path, data):
    tmp = f"{path}.{os.getpid()}.{threading.get_ident()}.tmp"   # загрузка и автообновление — не в один файл
    try:
        with open(tmp, "w", encoding="utf-8") as fh:
            json.dump(data, fh, ensure_ascii=False)
        os.replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)


def save(data, filename, path):
    ref = parse(data, filename)
    _write_json(path, ref)
    return ref


_summary = {"key": None, "value": None}


def summary(path):
    """Файл, число позиций и единицы справочника — без чтения всего справочника на каждый показ настроек."""
    try:
        s = os.stat(path)
    except OSError:
        return None
    key = (s.st_mtime_ns, s.st_size, s.st_ino)
    if _summary["key"] != key:
        ref = load(path)
        _summary.update(key=key, value=None if not ref else {
            "file": ref.get("file"), "items": len(ref.get("items") or {}), "units": ref.get("units") or {},
            "mtime": s.st_mtime})
    return _summary["value"]


def load(path):
    try:
        with open(path, encoding="utf-8") as fh:
            return json.load(fh)
    except (OSError, ValueError):
        return None


def _lev(a, b, cap=2):
    if abs(len(a) - len(b)) > cap:
        return cap + 1
    prev = list(range(len(b) + 1))
    for i, ca in enumerate(a, 1):
        cur = [i] + [0] * len(b)
        best = cur[0]
        for j, cb in enumerate(b, 1):
            cur[j] = min(prev[j] + 1, cur[j - 1] + 1, prev[j - 1] + (ca != cb))
            best = min(best, cur[j])
        if best > cap:
            return cap + 1
        prev = cur
    return prev[-1]


def _rows(dump):
    """Строки страниц, склеенные по высоте: ячейки одной строки таблицы → одна строка."""
    for pg in dump.get("lines", []):
        groups = []
        for ln in sorted(pg.get("lines", []), key=lambda x: x.get("y", 0)):
            if groups and abs(groups[-1][0] - ln.get("y", 0)) <= 2.5:
                groups[-1][1].append(ln["text"])
            else:
                groups.append([ln.get("y", 0), [ln["text"]]])
        for _, texts in groups:
            yield pg.get("pdf_page", pg.get("page", 1)), "  ".join(t.strip() for t in texts)


def _fmt(v):
    s = f"{v:.6f}".rstrip("0").rstrip(".")
    return s.replace(".", ",")


def _cell_num(w):
    w = w.strip().strip(";:")
    return num(w) if re.fullmatch(r"\d+(?:[.,]\d+)?", w) else None


def _page_rows(page, tol=3.0):
    """Слова страницы, сгруппированные в строки по высоте: [(y, [(x0, x1, текст), …]), …]."""
    rows = []
    for x0, y0, x1, y1, w, *_ in sorted(page.get_text("words"), key=lambda r: ((r[1] + r[3]) / 2, r[0])):
        yc = (y0 + y1) / 2
        if rows and abs(rows[-1][0] - yc) <= tol:
            rows[-1][1].append((x0, x1, w))
        else:
            rows.append([yc, [(x0, x1, w)]])
    for r in rows:
        r[1].sort()
    return rows


def _param_cols(words, params, units):
    """Колонки характеристик в строке заголовка: название («Мощность», «Напор номинальный») или уникальная единица
    («кВт»). → {характеристика: x центра}."""
    low = [w.lower().strip(",.:;()[]") for _, _, w in words]
    cols = {}
    for name in params:
        toks = [t for t in re.split(r"[\s,]+", name.lower()) if t][:2]
        if not toks:
            continue
        for i, w in enumerate(low):
            if len(toks[0]) >= 3 and w.startswith(toks[0][:6]) and (len(toks) == 1 or (i + 1 < len(low) and low[i + 1].startswith(toks[1][:4]))):
                j = i + (len(toks) > 1)
                cols[name] = (words[i][0] + words[j][1]) / 2
                break
    for name, unit in units.items():
        if name in cols:
            continue
        u = norm_unit(unit)
        hits = [i for i, w in enumerate(low) if norm_unit(w) == u]
        if len(hits) == 1:
            cols[name] = (words[hits[0]][0] + words[hits[0]][1]) / 2
    xs = list(cols.values())                           # одна колонка — одна характеристика
    return {n: x for n, x in cols.items() if xs.count(x) == 1}


def _nearest(words, x, maxd):
    best = None
    for x0, x1, w in words:
        v = _cell_num(w)
        if v is None:
            continue
        d = abs((x0 + x1) / 2 - x)
        if d <= maxd and (best is None or d < best[0]):
            best = (d, v)
    return best[1] if best else None


def check_tables(pdf_path, ref, st):
    """Сверка таблиц характеристик по колонкам — значения без единиц в ячейках, единица — в заголовке.
    Два вида таблиц: «строка — изделие» (заголовок с характеристиками сверху) и «колонка — изделие»
    (строка с артикулами сверху, характеристики — по строкам)."""
    import pymupdf
    items, units = ref.get("items") or {}, ref.get("units") or {}
    if not items or not units:
        return []
    tol = float(st.get("products_tolerance") if st.get("products_tolerance") is not None else 0.5) / 100
    out = []

    def cmp(page, art, name, val, row_text):
        p = items[art]["params"].get(name)
        if p is None or val is None or abs(val - p["value"]) <= max(abs(p["value"]) * tol, 1e-9):
            return
        out.append({"kind": "table", "page": page, "article": items[art]["article"], "suggest": f"{p['raw']} {p['unit']}",
                    "text": row_text[:200], "title": f"Расхождение в таблице со справочником: {name}",
                    "detail": f"{items[art]['article']}: в таблице {_fmt(val)} {p['unit']}, в справочнике {p['raw']} {p['unit']}"})

    with pymupdf.open(pdf_path) as doc:
        for pno, page in enumerate(doc, 1):
            rows = _page_rows(page)
            cols, gap = None, 40.0
            art_cols = None
            last_y = 0
            for y, words in rows:
                text = " ".join(w for _, _, w in words)
                if y - last_y > 60:                    # большой разрыв по высоте — таблица закончилась
                    cols = art_cols = None
                arts = [(norm(w), (x0 + x1) / 2) for x0, x1, w in words if norm(w) in items]
                # вид 2: строка с несколькими артикулами — шапка таблицы «колонка — изделие»
                if len(arts) >= 2:
                    art_cols, last_y = arts, y
                    xs = sorted(x for _, x in arts)
                    agap = min((b - a for a, b in zip(xs, xs[1:])), default=60) / 2
                    cols = None
                    continue
                hc = _param_cols(words, units.keys(), units)
                if len(hc) >= 2 or (hc and re.search(r"(?i)артикул|модель|обозначение|тип", text)):
                    cols, art_cols, last_y = hc, None, y  # вид 1: шапка «строка — изделие»
                    xs = sorted(cols.values())
                    gap = min((b - a for a, b in zip(xs, xs[1:])), default=80) / 2
                    continue
                if art_cols and len(hc) == 1 and words and (words[0][0] < min(x for _, x in art_cols)):
                    name = next(iter(hc))              # строка характеристики под шапкой с артикулами
                    for art, x in art_cols:
                        cmp(pno, art, name, _nearest(words, x, agap), text)
                    last_y = y
                    continue
                if cols and len({a for a, _ in arts}) == 1:
                    art = arts[0][0]
                    last_y = y
                    for name, x in cols.items():
                        cmp(pno, art, name, _nearest([w for w in words if norm(w[2]) != art], x, gap), text)
            if len(out) >= 300:
                break
    return out


def check(dump, ref, st, pdf_path=None):
    items = ref.get("items") or {}
    if not items:
        return None
    tol = float(st.get("products_tolerance", 0.5) if st.get("products_tolerance") is not None else 0.5) / 100
    by_prefix = defaultdict(list)                      # индекс: (2 первых знака, длина) → артикулы
    for k in items:
        by_prefix[(k[:2], len(k))].append(k)
    series = Counter(m.group(0) for k in items for m in [re.match(r"[A-Z]{2,}", k)] if m)
    series = {p for p, n in series.items() if n >= 3}
    found, issues, seen = set(), [], set()
    unit_re = {}
    unit_count = Counter(norm_unit(u) for u in (ref.get("units") or {}).values())
    for name, unit in (ref.get("units") or {}).items():
        u = r"\s?".join(re.escape(ch) for ch in norm_unit(unit))
        unit_re[name] = re.compile(r"(?<![\w.,])(\d+(?:[.,]\d+)?)\s*" + u + r"(?![\w])", re.I)
    for page, row in _rows(dump):
        known = []
        for m in TOKEN.finditer(row):
            tok = m.group(0)
            k = norm(tok)
            if k in items:
                found.add(k)
                known.append(k)
                continue
            if (k, page) in seen or len(k) < 4 or not any(ch.isupper() for ch in tok):
                continue                                # «м3/ч», «мм2» и т. п. — единицы, а не артикулы
            cap = 2 if len(k) >= 8 else 1
            cand = [a for n in range(len(k) - cap, len(k) + cap + 1) for a in by_prefix.get((k[:2], n), [])][:3000]
            close = sorted((d, abs(len(a) - len(k)), a) for a in cand for d in [_lev(k, a, cap)] if d <= cap)
            close = [a for _, _, a in close]
            if close:
                seen.add((k, page))
                issues.append({"kind": "typo", "page": page, "article": tok, "suggest": items[close[0]]["article"],
                               "text": row[:200], "title": "Возможна опечатка в артикуле",
                               "detail": f"«{tok}» нет в справочнике; похожий артикул: «{items[close[0]]['article']}»"})
            elif st.get("products_unknown", True) and any(k.startswith(p) and len(k) > len(p) and not k[len(p)].isalpha()
                                                          for p in series):
                seen.add((k, page))
                issues.append({"kind": "unknown", "page": page, "article": tok, "suggest": "", "text": row[:200],
                               "title": "Артикул не найден в справочнике",
                               "detail": f"«{tok}» похож на артикул серии из справочника, но в справочнике его нет"})
        if len(set(known)) != 1 or not unit_re:
            continue
        it = items[known[0]]
        low = row.replace("³", "3").replace("²", "2").lower()
        for name, rx in unit_re.items():
            p = it["params"].get(name)
            if not p:
                continue
            if unit_count[norm_unit(p["unit"])] > 1 and name.lower() not in low:
                continue                               # «Высота, мм» и «Ширина, мм»: без названия в строке не различить
            vals = [num(v) for v in rx.findall(low)]
            vals = [v for v in vals if v is not None]
            if not vals or any(abs(v - p["value"]) <= max(abs(p["value"]) * tol, 1e-9) for v in vals):
                continue
            issues.append({"kind": "param", "page": page, "article": it["article"], "suggest": f"{p['raw']} {p['unit']}",
                           "text": row[:200], "title": f"Расхождение со справочником: {name}",
                           "detail": f"{it['article']}: в документе {str(round(vals[0], 6)).rstrip('0').rstrip('.').replace('.', ',')} {p['unit']}, в справочнике {p['raw']} {p['unit']}"})
        if len(issues) >= 500:
            break
    if pdf_path:
        try:
            issues += check_tables(pdf_path, ref, st)
        except Exception:  # noqa — сверка таблиц необязательна
            pass
    grouped = {}                                       # одно замечание на артикул и расхождение — со списком страниц
    for it in issues:
        g = grouped.setdefault((it["kind"], it["article"], it["detail"]), dict(it, pages=[]))
        if it["page"] not in g["pages"]:
            g["pages"].append(it["page"])
    issues = list(grouped.values())
    return {"loaded": len(items), "file": ref.get("file", ""), "found": len(found),
            "found_list": sorted(items[k]["article"] for k in found)[:300], "issues": issues[:500],
            "ok": not issues}


# ------------------------------------------------------------------ автообновление по адресу
def sync_status_path():
    return os.path.join(os.path.dirname(PATH()), ".products_sync.json")


def sync_status():
    try:
        with open(sync_status_path(), encoding="utf-8") as fh:
            return json.load(fh)
    except (OSError, ValueError):
        return {}


def sync(st, max_bytes=50 * 1024 * 1024):
    """Скачивает справочник по адресу из настроек и загружает его. → итог словарем (ошибка — в "error")."""
    import datetime as dt
    import httpx
    url = (st.get("products_sync_url") or "").strip()
    res = {"time": dt.datetime.now().isoformat(timespec="seconds"), "url": url}
    try:
        if not re.match(r"^https?://", url):
            raise ValueError("адрес выгрузки должен начинаться с http:// или https://")
        headers = {"User-Agent": "Korrektor"}
        if st.get("products_sync_auth"):
            headers["Authorization"] = st["products_sync_auth"].strip()
        buf = bytearray()
        deadline = time.time() + SYNC_TIMEOUT        # тайм-аут httpx — на каждое чтение; медленный сервер, отдающий
        with httpx.stream("GET", url, headers=headers, timeout=60, follow_redirects=True) as r:   # по байту, держал бы
            if r.status_code != 200:                                                               # поток часами
                raise ValueError(f"сервер ответил {r.status_code}")
            for chunk in r.iter_bytes():
                buf += chunk
                if len(buf) > max_bytes:
                    raise ValueError("выгрузка больше 50 МБ")
                if time.time() > deadline:
                    raise ValueError(f"выгрузка не скачалась за {SYNC_TIMEOUT // 60} мин")
            ctype = r.headers.get("content-type", "")
        name = os.path.basename(url.split("?")[0]) or "products.csv"
        if not name.lower().endswith((".csv", ".txt", ".xlsx", ".xlsm")):
            name += ".xlsx" if ("spreadsheet" in ctype or bytes(buf[:2]) == b"PK") else ".csv"
        ref = save(bytes(buf), name, PATH())
        res.update(ok=True, items=len(ref["items"]), units=ref["units"], file=name)
    except Exception as e:  # noqa
        res.update(ok=False, error=str(e)[:300])
    try:
        _write_json(sync_status_path(), res)
    except OSError:
        pass
    return res


SYNC_TIMEOUT = 600


def sync_loop(load_settings, log):
    import datetime as dt
    tried = None                                     # и при сбое записи статуса — не чаще раза в день
    while True:
        try:
            st = load_settings()
            now = dt.datetime.now()
            last = sync_status().get("time", "")[:10]
            today = now.date().isoformat()
            from .. import core
            if core.PAUSED.is_set():                 # восстановление из копии
                time.sleep(60)
                continue
            if st.get("products_sync_url") and now.hour >= int(st.get("products_sync_hour", 6) or 0) \
                    and last != today and tried != today:
                tried = today
                r = sync(st)
                if not r.get("ok"):
                    log.warning("автообновление справочника: %s", r.get("error"))
        except Exception:  # noqa
            log.exception("автообновление справочника")
        time.sleep(900)
