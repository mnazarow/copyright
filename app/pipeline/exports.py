# -*- coding: utf-8 -*-
"""Дополнительные форматы отчета: XLSX и CSV (основной DOCX собирает build.py из мастер-промпта)."""
import re
import csv

COLS = ["№", "Стр., пункт", "Тип", "Как в документе", "Как правильно", "Пояснение"]


CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f\ud800-\udfff\ufffe\uffff]")


def safe(v):
    """Защита от «формульной инъекции»: текст из PDF, начинающийся с = + - @, Excel не должен считать формулой.
    Управляющие знаки удаляются: Excel и Word их не принимают."""
    if isinstance(v, str):
        v = CTRL.sub("", v)
        return "'" + v if v[:1] in ("=", "+", "-", "@", "\t", "\r") else v
    return v


def rows(findings):
    for i, f in enumerate(findings, 1):
        yield [safe(x) for x in [i, f"стр. {f['page']}, {f['clause']}", f["type"], f["src"].replace(" ", " "),
               f["fix"].replace(" ", " "), f["note"]]]


def to_csv(findings, path):
    with open(path, "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow(COLS)
        for r in rows(findings):
            w.writerow(r)


def to_xlsx(findings, meta, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo
    wb = Workbook()
    ws = wb.active
    ws.title = "Замечания"
    ws.append(COLS)
    for r in rows(findings):
        ws.append(r)
    widths = [6, 22, 22, 50, 50, 70]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0F4382")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    if findings:
        t = Table(displayName="Zamechaniya", ref=f"A1:F{len(findings) + 1}")
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws.add_table(t)
    ws.freeze_panes = "A2"
    info = wb.create_sheet("Документ")
    for line in meta.get("header_lines", []) + [meta.get("totals", ""), "", "Итог:", meta.get("summary", "")]:
        info.append([safe(line)])
    info.column_dimensions["A"].width = 140
    wb.save(path)
