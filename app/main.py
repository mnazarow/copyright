# -*- coding: utf-8 -*-
"""Веб-сервис «Корректор»: загрузка PDF, отчет о корректорской вычитке, настройки, документация."""
import collections
import contextlib
import datetime as dt
import io
import json
import os
import re
import asyncio
import hmac
from html import escape as html_escape
import secrets
import time
import zipfile
import zlib
from urllib.parse import quote, urlsplit

import markdown
from markdown.extensions.toc import slugify_unicode
from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, RedirectResponse, Response
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from starlette.concurrency import run_in_threadpool

from . import admin, analytics, oidc, totp, users
from . import batch as batchmod
from . import core
from .pipeline import office, run
from .pipeline.common import AUTHOR, TYPES
from .settings_schema import DEFAULTS, GROUPS, PROFILES, S, SCHEMA, profile_overrides

@contextlib.asynccontextmanager
async def lifespan(_app):
    core.start_workers()
    yield


app = FastAPI(title="Корректор", version=core.VERSION, docs_url="/api/docs", redoc_url=None, lifespan=lifespan)
app.mount("/static", StaticFiles(directory=os.path.join(core.APP_DIR, "static")), name="static")
DOCS_DIR = os.path.join(core.ROOT, "docs")
app.mount("/docs-files", StaticFiles(directory=DOCS_DIR), name="docs-files")
tpl = Jinja2Templates(directory=os.path.join(core.APP_DIR, "templates"))
tpl.env.globals.update(version=core.VERSION, TYPES=TYPES, AUTHOR=AUTHOR)

DOC_PAGES = [
    ("index", "Обзор и быстрый старт"),
    ("methodology", "Методология вычитки"),
    ("usage", "Работа с сервисом"),
    ("analytics", "Аналитика отчетов"),
    ("trends", "Тренды"),
    ("users", "Пользователи, роли и журнал"),
    ("integrations", "Уведомления и интеграции"),
    ("standards", "Стандарт оформления и справочник"),
    ("settings", "Справочник настроек"),
    ("installation", "Установка, обновление, удаление"),
    ("administration", "Администрирование и безопасность"),
    ("api", "REST API и интеграция"),
    ("faq", "Вопросы и неполадки"),
]
PUBLIC = ("/login", "/static", "/api/health")
SETUP_2FA_OK = ("/account", "/logout", "/static", "/login")      # пока обязательный второй фактор не подключен


# ------------------------------------------------------------------ доступ
def authed(request: Request):
    """→ (пользователь, способ): способ 'cookie' — браузер, 'token' — токен REST API (только /api/…), 'open' — вход отключен."""
    st = core.load_settings()
    if not st.get("auth_required", True) or not core.ADMIN_PASSWORD:
        return users.get(users.BUILTIN), "open"
    tok = core.check_token(request.cookies.get("korrektor", ""))
    if tok:
        u = users.get(tok[0])
        if u and u.get("active") and hmac.compare_digest(tok[1].encode(), users.session_stamp(u).encode()):
            return u, "cookie"
    tok = st.get("api_token")
    auth = request.headers.get("authorization", "")
    if tok and request.url.path.startswith("/api/") and hmac.compare_digest(auth.encode(), f"Bearer {tok}".encode()):
        return {"id": -1, "login": "api", "name": "REST API", "role": "corrector", "source": "token", "active": 1,
                "email": "", "department": ""}, "token"
    return None, None


def same_origin(request: Request):
    """Защита от межсайтовых запросов: POST из браузера должен прийти со страницы этого же сервиса."""
    src = request.headers.get("origin") or request.headers.get("referer")
    if not src:
        return True                                   # не браузер (curl, скрипт) — браузеры всегда шлют Origin
    host = request.headers.get("x-forwarded-host") or request.headers.get("host", "")
    return urlsplit(src).netloc == host


@app.middleware("http")
async def guard(request: Request, call_next):
    path = request.url.path
    request.state.user = None
    st_ = core.load_settings()
    if path != "/api/health" and (st_.get("allowed_networks") or st_.get("admin_networks")):
        ip = users.client_ip(request)
        if not users.ip_allowed(ip, st_.get("allowed_networks")) or (
                path.startswith(users.ADMIN_PATHS) and not users.ip_allowed(ip, st_.get("admin_networks"))):
            core.log.warning("доступ запрещен: %s %s", ip, path)
            msg = f"Доступ с адреса {ip} запрещен настройками сервиса (разрешенные сети). Обратитесь к администратору."
            if path.startswith("/api/"):
                return JSONResponse({"error": msg}, status_code=403)
            return HTMLResponse(f"<!doctype html><meta charset='utf-8'><title>Доступ запрещен</title>"
                                f"<p style='font:16px Arial;margin:40px'>{html_escape(msg)}</p>", status_code=403)
    if path.startswith(PUBLIC) or path == "/metrics":
        return await call_next(request)
    user, how = authed(request)
    if not user:
        if path.startswith("/api/"):
            return JSONResponse({"error": "требуется авторизация"}, status_code=401)
        return RedirectResponse("/login?next=" + quote(path), status_code=303)
    request.state.user = user
    if how == "cookie" and not path.startswith(SETUP_2FA_OK) and totp.required(user, core.load_settings()) \
            and not totp.enabled(user["login"]):
        # обязательный второй фактор включили, когда сеанс уже был открыт: сначала подключить
        if path.startswith("/api/") or request.method != "GET":
            return JSONResponse({"error": "подключите двухфакторный вход в профиле"}, status_code=403)
        return RedirectResponse("/account/2fa?required=1", status_code=303)
    if core.PAUSED.is_set() and request.method not in ("GET", "HEAD", "OPTIONS"):
        return JSONResponse({"error": "идет восстановление из резервной копии — повторите через минуту"}, status_code=503)
    if request.method not in ("GET", "HEAD", "OPTIONS") and how != "token" and not same_origin(request):
        return JSONResponse({"error": "запрос с чужой страницы отклонен"}, status_code=403)
    # номера заданий — только в каноническом виде (FastAPI принял бы и «+1», «1.0», « 1» — обход проверок)
    m = users.JOB_PATH.match(path)
    ids = [m.group(1)] if m else []
    bad_id = False
    if path.startswith(("/compare", "/api/compare")):
        vals = [(k, v) for k, v in request.query_params.multi_items() if k in ("a", "b", "job", "p")]
        bad_id = any(not users.ID_RE.fullmatch(v) for _, v in vals)
        ids = [v for k, v in vals if k != "p"]
    if bad_id or any(not users.ID_RE.fullmatch(i) for i in ids):
        return JSONResponse({"error": "неверный номер задания"}, status_code=404)
    if ids and users.own_only(user) is not None and not all(users.job_visible(user, int(i)) for i in ids):
        if path.startswith("/api/"):
            return JSONResponse({"error": "задание не найдено"}, status_code=404)
        return HTMLResponse(tpl.get_template("denied.html").render(ctx(request, "denied", msg="Это задание загрузил другой сотрудник")),
                            status_code=404)
    if not users.path_allowed(user, request.method, path):
        msg = f"Недостаточно прав: роль «{users.ROLES.get(user['role'], user['role'])}»"
        if path.startswith("/api/") or request.method != "GET":
            return JSONResponse({"error": msg}, status_code=403)
        return HTMLResponse(tpl.get_template("denied.html").render(ctx(request, "denied", msg=msg)), status_code=403)
    return await call_next(request)


def safe_next(nxt):
    """Только относительный адрес внутри сервиса (защита от открытого перенаправления)."""
    nxt = (nxt or "/").strip()
    if not nxt.startswith("/") or nxt.startswith("//") or nxt.startswith("/\\") or any(ord(ch) < 32 or ch == "\x7f" for ch in nxt):
        return "/"
    return nxt


@app.get("/login", response_class=HTMLResponse)
def login_page(request: Request, next: str = "/", error: str = ""):
    st = core.load_settings()
    return tpl.TemplateResponse(request, "login.html", {"next": next, "error": error, "page": "login",
                                                        "ldap": st.get("ldap_enabled"),
                                                        "oidc": oidc.enabled(st) and (st.get("oidc_button") or "Единый вход")})


_fails = {}                                           # ключ → [попыток, время блокировки]
_fails_lock = __import__("threading").Lock()


def _take_attempt(keys_limits):
    """Атомарно занимает попытку входа для всех ключей до проверки пароля или кода: параллельные запросы
    не обходят ограничение. → False, если хотя бы один ключ заблокирован."""
    now = time.time()
    with _fails_lock:
        if len(_fails) > 20000:                        # чистим только истекшие записи
            for k in [k for k, v in _fails.items() if v[1] < now]:   # действующие блокировки сохраняются
                _fails.pop(k, None)
        if len(_fails) > 50000:                        # поток незаблокированных записей — память не растет без предела
            for k in list(_fails)[:25000]:
                if not _fails[k][1]:
                    _fails.pop(k, None)
        for k, lim in keys_limits:
            cnt, until = _fails.get(k, [0, 0])
            if until and until <= now:
                _fails[k] = [0, 0]
            elif until > now or cnt >= lim:
                return False
        for k, lim in keys_limits:
            cnt = _fails.get(k, [0, 0])[0] + 1
            _fails[k] = [cnt, now + 300 if cnt >= lim else 0]
        return True


def _attempt_ok(key, ip=None):
    """Успешный вход: счетчик логина сбрасывается, а занятая попытка адреса возвращается."""
    with _fails_lock:
        _fails.pop(key, None)
        if ip and ip in _fails and not _fails[ip][1]:
            _fails[ip][0] = max(0, _fails[ip][0] - 1)


def _attempts(key):
    return _fails.get(key, [0, 0])[0]


@app.post("/login")
async def login(request: Request, password: str = Form(""), next: str = Form("/"), login: str = Form("")):
    if not same_origin(request):                       # вход с чужой страницы (подмена учетной записи) — отклоняется
        return RedirectResponse("/login?error=1", status_code=303)
    ip = users.client_ip(request) or "?"
    key = f"{ip}|{(login or 'admin').strip().lower()[:64]}"
    # 5 ошибок подряд для логина — пауза 5 минут; 30 ошибок с адреса — пауза для адреса
    if not _take_attempt([(key, 5), (ip, 30)]):
        return RedirectResponse("/login?error=locked&next=" + quote(safe_next(next)), status_code=303)
    st = core.load_settings()
    user, why = await run_in_threadpool(users.authenticate, login, password, st)
    if not user:
        request.state.user = {"login": (login or "admin")[:64]}
        users.audit(request, "login_failed", login or "admin", why or "")
        await asyncio.sleep(1)
        err = "locked" if _attempts(key) >= 5 else ("ldap" if why == "ldap" else "disabled" if why == "disabled" else "1")
        return RedirectResponse("/login?error=" + err + "&next=" + quote(safe_next(next)), status_code=303)
    _attempt_ok(key, ip)
    if totp.enabled(user["login"]) or totp.required(user, st):
        # второй шаг: код из приложения. Промежуточный пропуск подписан с другой «солью» — сеансом он не является
        r = RedirectResponse(("/login/2fa" if totp.enabled(user["login"]) else "/login/2fa-setup")
                             + "?next=" + quote(safe_next(next)), status_code=303)
        r.set_cookie("korrektor_2fa", core.make_token(user["login"], ttl=600, stamp="2fa:" + users.session_stamp(user)),
                     httponly=True, samesite="lax", max_age=600, secure=_https(request))
        return r
    return _start_session(request, user, st, next)


def _https(request):
    return request.url.scheme == "https" or request.headers.get("x-forwarded-proto") == "https"


def _start_session(request, user, st, nxt, detail="", resp=None):
    users.touch(user["login"])
    request.state.user = user
    users.audit(request, "login", user["login"], detail)
    hours = int(st.get("session_hours") or 336)
    r = resp or RedirectResponse(safe_next(nxt), status_code=303)
    r.set_cookie("korrektor", core.make_token(user["login"], ttl=hours * 3600, stamp=users.session_stamp(user)),
                 httponly=True, samesite="lax", max_age=hours * 3600, secure=_https(request))
    r.delete_cookie("korrektor_2fa")
    r.delete_cookie("korrektor_2fa_new")
    return r


def _pending(request):
    """Пользователь, прошедший первый шаг входа (пароль), или None."""
    tok = core.check_token(request.cookies.get("korrektor_2fa", ""))
    if not tok:
        return None
    u = users.get(tok[0])
    if not u or not u.get("active") or not hmac.compare_digest(tok[1].encode(), ("2fa:" + users.session_stamp(u)).encode()):
        return None
    return u


def _tpl2fa(request, **kw):
    return tpl.TemplateResponse(request, "login2fa.html", {"page": "login", **kw})


@app.get("/login/2fa", response_class=HTMLResponse)
def login_2fa_page(request: Request, next: str = "/", error: str = ""):
    u = _pending(request)
    if not u:
        return RedirectResponse("/login?error=expired", status_code=303)
    return _tpl2fa(request, mode="code", next=safe_next(next), error=error, login=u["login"])


@app.post("/login/2fa")
async def login_2fa(request: Request, code: str = Form(""), next: str = Form("/")):
    if not same_origin(request):
        return RedirectResponse("/login?error=expired", status_code=303)
    u = _pending(request)
    if not u:
        return RedirectResponse("/login?error=expired", status_code=303)
    key = "2fa|" + u["login"].lower()
    if not _take_attempt([(key, 5)]):                  # 5 неверных кодов — пауза 5 минут (атомарно, до проверки)
        r = RedirectResponse("/login?error=locked", status_code=303)
        r.delete_cookie("korrektor_2fa")
        return r
    how = await run_in_threadpool(totp.verify, u["login"], code)
    if not how:
        request.state.user = {"login": u["login"]}
        users.audit(request, "login_failed", u["login"], "неверный код второго фактора")
        await asyncio.sleep(1)
        if _attempts(key) >= 5:
            r = RedirectResponse("/login?error=locked", status_code=303)
            r.delete_cookie("korrektor_2fa")
            return r
        return RedirectResponse("/login/2fa?error=1&next=" + quote(safe_next(next)), status_code=303)
    _attempt_ok(key)
    left = len((totp.get(u["login"]) or {}).get("recovery") or [])
    detail = "двухфакторный вход" if how == "totp" else f"резервный код (осталось {left})"
    return _start_session(request, u, core.load_settings(), next, detail)


@app.get("/login/2fa-setup", response_class=HTMLResponse)
def login_2fa_setup_page(request: Request, next: str = "/", error: str = ""):
    """Обязательное подключение второго фактора при входе (настройка «Двухфакторный вход обязателен для»)."""
    u = _pending(request)
    if not u or totp.enabled(u["login"]):
        return RedirectResponse("/login?error=expired", status_code=303)
    new = oidc._unsign(request.cookies.get("korrektor_2fa_new", "")) or {}
    secret = new.get("s") if new.get("u") == u["login"] and new.get("t") == "totp" else None
    fresh = not secret
    secret = secret or totp.new_secret()
    resp = _tpl2fa(request, mode="setup", next=safe_next(next), error=error, login=u["login"], secret=secret,
                   action="/login/2fa-setup", uri=totp.uri(secret, u["login"]))
    if fresh:
        resp.set_cookie("korrektor_2fa_new", oidc._sign({"t": "totp", "s": secret, "u": u["login"], "exp": time.time() + 600}),
                        httponly=True, samesite="lax", max_age=600, secure=_https(request))
    return resp


@app.post("/login/2fa-setup")
async def login_2fa_setup(request: Request, code: str = Form(""), next: str = Form("/")):
    u = _pending(request)
    if not u or totp.enabled(u["login"]):
        return RedirectResponse("/login?error=expired", status_code=303)
    new = oidc._unsign(request.cookies.get("korrektor_2fa_new", "")) or {}
    if new.get("u") != u["login"] or not new.get("s") or new.get("t") != "totp":
        return RedirectResponse("/login/2fa-setup?next=" + quote(safe_next(next)), status_code=303)
    step = totp.check_code(new["s"], code)
    if not step:
        await asyncio.sleep(1)
        return RedirectResponse("/login/2fa-setup?error=1&next=" + quote(safe_next(next)), status_code=303)
    codes = totp.enable(u["login"], new["s"], step)
    u = users.get(u["login"])                          # отпечаток сеанса изменился вместе со вторым фактором
    request.state.user = u
    users.audit(request, "totp", u["login"], "подключен при входе")
    page = _tpl2fa(request, mode="codes", codes=codes, next=safe_next(next), login=u["login"])
    return _start_session(request, u, core.load_settings(), next, "двухфакторный вход подключен", resp=page)


@app.get("/login/oidc")
async def login_oidc(request: Request, next: str = "/"):
    st = core.load_settings()
    if not oidc.enabled(st):
        raise HTTPException(404)
    try:
        url, cookie = await run_in_threadpool(oidc.start, st, request, safe_next(next))
    except Exception as e:  # noqa
        core.log.warning("OIDC: %s", e)
        return RedirectResponse("/login?error=oidc", status_code=303)
    r = RedirectResponse(url, status_code=303)
    r.set_cookie("korrektor_oidc", cookie, httponly=True, samesite="lax", max_age=600,
                 secure=request.url.scheme == "https" or request.headers.get("x-forwarded-proto") == "https")
    return r


@app.get("/login/oidc/callback")
async def login_oidc_callback(request: Request, code: str = "", state: str = "", error: str = "",
                              error_description: str = ""):
    st = core.load_settings()
    if not oidc.enabled(st):
        raise HTTPException(404)
    if error or not code:
        users.audit(request, "login_failed", "oidc", (error + " " + error_description)[:300])
        return RedirectResponse("/login?error=oidc", status_code=303)
    try:
        user, nxt = await run_in_threadpool(oidc.finish, st, request, code, state, request.cookies.get("korrektor_oidc"))
    except Exception as e:  # noqa
        core.log.warning("OIDC: %s", e)
        users.audit(request, "login_failed", "oidc", str(e)[:300])
        return RedirectResponse("/login?error=oidc", status_code=303)
    r = _start_session(request, user, st, nxt, "OpenID Connect")
    r.delete_cookie("korrektor_oidc")
    return r


@app.get("/logout")
def logout(request: Request):
    users.audit(request, "logout")
    r = RedirectResponse("/login", status_code=303)
    r.delete_cookie("korrektor")
    return r


# ------------------------------------------------------------------ страницы
def ctx(request, page, **kw):
    st = core.load_settings()
    user = getattr(request.state, "user", None)
    kw.update(request=request, page=page, st=st, user=user, ROLES=users.ROLES,
              can=lambda a: users.can(user, a),
              auth_on=bool(core.ADMIN_PASSWORD) and st.get("auth_required", True),
              ai_on=st.get("ai_provider") not in (None, "", "none"))
    return kw


@app.get("/", response_class=HTMLResponse)
def index(request: Request):
    own = users.own_only(request.state.user)
    jobs = core.list_jobs(6, owner=own)
    authors = []
    for j in core.list_jobs(2000, owner=own):
        if j.get("author") and j["author"] not in authors:
            authors.append(j["author"])
    authors = authors[:50]
    return tpl.TemplateResponse(request, "index.html", ctx(request, "index", jobs=jobs, authors=authors,
                                                           profiles=SCHEMA["profile"]["options"],
                                                           modes=SCHEMA["mode"]["options"]))


@app.post("/upload")
async def upload(request: Request, files: list[UploadFile] = File(...), profile: str = Form(""),
                 mode: str = Form(""), note: str = Form(""), author: str = Form(""), department: str = Form(""),
                 doc_type: str = Form("")):
    jobs = await _create_jobs(request, files, profile, mode, note, author, department, doc_type)
    if not jobs:
        raise HTTPException(400, "Не найдено ни одного PDF или документа Word")
    if len(jobs) == 1:
        return RedirectResponse(f"/job/{jobs[0]}", status_code=303)
    return RedirectResponse(f"/batch/{core.get_job(jobs[0])['batch']}", status_code=303)


async def _create_jobs(request, files, profile="", mode="", note="", author="", department="", doc_type=""):
    st = core.load_settings()
    user = getattr(request.state, "user", None) or {}
    author = (author or "").strip()[:120] or (user.get("name") or "" if user.get("role") == "author" else "")
    department = (department or "").strip()[:120] or (user.get("department") or "")
    if st.get("upload_author_required") and not author:
        raise HTTPException(400, "Укажите автора документа (настройка «Требовать автора документа при загрузке»)")
    limit = int(st.get("max_upload_mb") or 200) * 1024 * 1024
    over = {}
    if profile and profile in PROFILES:
        over.update(profile_overrides(profile))
        over = {k: v for k, v in over.items() if v != st.get(k)}
    if mode:
        over["mode"] = mode
    if note:
        over["report_extra_header"] = (st.get("report_extra_header") or "") + "\n" + note
    if st.get("department_scope") and user.get("role") != "admin" and user.get("source") != "token":
        department = (users.get(user.get("login", "")) or user).get("department") or ""   # документ — в отдел загрузившего
    meta = {"owner": user.get("login", ""), "author": author, "department": department,
            "doc_type": (doc_type or "").strip()[:80]}
    bid = core.new_batch_id()
    ids = []
    total = 0
    if len(files) > 200:
        raise HTTPException(413, "За одну загрузку — не больше 200 файлов")
    budget = {"bytes": limit * 3, "jobs": 200}             # на весь запрос, включая содержимое архивов
    try:
        for f in files:
            name = os.path.basename(f.filename or "document.pdf")
            if f.size is not None and f.size > limit:      # отказ до чтения файла в память
                raise HTTPException(413, f"Файл {name} больше {st.get('max_upload_mb')} МБ")
            data = await f.read(limit + 1)
            if len(data) > limit:
                raise HTTPException(413, f"Файл {name} больше {st.get('max_upload_mb')} МБ")
            total += len(data)
            if total > limit * 3:
                raise HTTPException(413, f"За одну загрузку — не больше {limit * 3 // 2**20} МБ")
            ids += await run_in_threadpool(_store_upload, bid, name, data, over, limit, meta,
                                           bool(st.get("office_convert", True)), budget, ids)
    except BaseException:
        for jid in ids:                                    # загрузка прервана — уже созданные задания убираются
            try:
                core.delete_job(jid)
            except Exception:  # noqa
                pass
        raise
    if ids:
        users.audit(request, "upload", ", ".join(map(str, ids)), {"files": len(files), "author": author,
                                                                   "department": department, "type": doc_type})
    core.kick()
    return ids


def _store_upload(bid, name, data, over, limit, meta=None, office_ok=True, budget=None, created=None):
    """Сохраняет PDF (или документ Word/ODT/RTF, а также файлы из ZIP) как задания. Выполняется в пуле потоков.
    budget — общий на запрос остаток распакованных байт и заданий; created — список, куда сразу пишутся номера
    (для отката при ошибке на середине)."""
    ids = created if created is not None else []
    start = len(ids)
    meta = meta or {}
    budget = budget if budget is not None else {"bytes": limit * 3, "jobs": 200}

    def add(fname, blob, ext="pdf"):
        budget["bytes"] -= len(blob)
        budget["jobs"] -= 1
        if budget["bytes"] < 0 or budget["jobs"] < 0:
            raise HTTPException(413, f"За одну загрузку — не больше {limit * 3 // 2**20} МБ после распаковки и "
                                     "не больше 200 документов")
        stem = re.sub(r"\.(pdf|docx?|odt|rtf)$", "", fname, flags=re.I)
        jid = core.create_job(bid, stem, fname, len(blob), over, **meta)
        d = core.job_dir(jid)
        os.makedirs(d, exist_ok=True)
        ids.append(jid)
        with open(os.path.join(d, "source." + ext), "wb") as fh:
            fh.write(blob)
        core.update_job(jid, status="queued", step="В очереди" if ext == "pdf" else "В очереди (будет преобразован в PDF)")

    def take(fname, blob):
        if blob[:5] == b"%PDF-":
            add(fname, blob)
        elif office_ok and office.is_office(fname, blob):
            add(fname, blob, os.path.splitext(fname.lower())[1][1:])

    if name.lower().endswith(".zip"):
        try:
            z = zipfile.ZipFile(io.BytesIO(data))
        except zipfile.BadZipFile:
            raise HTTPException(400, f"Файл {name} — поврежденный ZIP-архив")
        with z:
            docs = [zi for zi in z.infolist() if not zi.is_dir() and not os.path.basename(zi.filename).startswith(("~$", "."))
                    and (zi.filename.lower().endswith(".pdf") or (office_ok and office.is_office(zi.filename)))]
            if sum(zi.file_size for zi in docs) > limit * 3 or len(docs) > 200:      # защита от «zip-бомб»
                raise HTTPException(413, f"Архив {name}: после распаковки больше {limit * 3 // 2**20} МБ "
                                         "или больше 200 файлов")
            for zi in docs:
                try:
                    blob = z.read(zi)
                except (zipfile.BadZipFile, zlib.error, RuntimeError):
                    raise HTTPException(400, f"Архив {name} поврежден или зашифрован ({zi.filename})")
                take(os.path.basename(zi.filename), blob)
    else:
        take(name, data)
    return ids[start:] if created is None else []


@app.get("/jobs", response_class=HTMLResponse)
def jobs_page(request: Request):
    return tpl.TemplateResponse(request, "jobs.html", ctx(request, "jobs", jobs=core.list_jobs(owner=users.own_only(request.state.user))))


@app.get("/batch/{bid}", response_class=HTMLResponse)
def batch_page(request: Request, bid: str):
    jobs = core.list_jobs(batch=bid, owner=users.own_only(request.state.user))
    if not jobs:
        raise HTTPException(404)
    return tpl.TemplateResponse(request, "batch.html", ctx(request, "jobs", bid=bid, jobs=jobs, res=batchmod.load(bid)))


@app.get("/job/{jid}", response_class=HTMLResponse)
def job_page(request: Request, jid: int):
    job = core.get_job(jid)
    if not job:
        raise HTTPException(404)
    state = core.load_state(jid)
    logtxt = ""
    lp = os.path.join(core.job_dir(jid), "job.log")
    if os.path.exists(lp):
        logtxt = open(lp, encoding="utf-8").read()[-6000:]
    bsize = len(core.list_jobs(batch=job["batch"])) if job.get("batch") else 1
    return tpl.TemplateResponse(request, "job.html", ctx(request, "jobs", job=job, state=state, log=logtxt, bsize=bsize))


@app.get("/job/{jid}/file/{kind}")
def job_file(jid: int, kind: str):
    job = core.get_job(jid)
    if not job:
        raise HTTPException(404)
    if kind == "source":
        src = office.find_source(core.job_dir(jid)) or os.path.join(core.job_dir(jid), "source.pdf")
        if not os.path.exists(src):
            raise HTTPException(404)
        return FileResponse(src, filename=job["filename"])
    state = core.load_state(jid)
    if not state:
        raise HTTPException(404, "Отчет еще не готов")
    st = core.load_settings()
    v = state.get("verify", {})
    fn = v.get("files", {}).get(kind)
    if kind == "docx" and not v.get("ok") and not st.get("report_allow_unverified"):
        raise HTTPException(403, "Отчет не прошел машинную проверку (п. 0.9). Скачивание непринятого отчета "
                                 "можно разрешить в настройках: «Отчет» → «Разрешать скачивание непринятого отчета».")
    if not fn:
        raise HTTPException(404, "Этот формат не создавался — включите его в настройках «Отчет» → «Форматы отчета»")
    path = os.path.join(core.job_dir(jid), fn)
    return FileResponse(path, filename=os.path.basename(fn))


@app.get("/job/{jid}/page/{n}.png")
def job_page_img(jid: int, n: int):
    p = os.path.join(core.job_dir(jid), "pages", f"p{n:03d}.png")
    if not os.path.exists(p):
        src = os.path.join(core.job_dir(jid), "source.pdf")    # после восстановления из копии без изображений
        if n < 1 or not os.path.exists(src):
            raise HTTPException(404)
        try:
            import pymupdf
            from .pipeline.common import safe_dpi
            with pymupdf.open(src) as d:
                if n > d.page_count:
                    raise HTTPException(404)
                os.makedirs(os.path.dirname(p), exist_ok=True)
                d[n - 1].get_pixmap(dpi=safe_dpi(d[n - 1], int(core.load_settings().get("page_images_dpi") or 72))).save(p + ".tmp.png")
                os.replace(p + ".tmp.png", p)
        except HTTPException:
            raise
        except Exception:  # noqa
            raise HTTPException(404)
    return FileResponse(p, media_type="image/png")


@app.get("/job/{jid}/checks.xlsx")
def job_checks_xlsx(jid: int):
    """Стандарт оформления и сверка со справочником продукции — отдельной таблицей."""
    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from .pipeline.exports import safe as _s

    def _safe(v):
        return _s(ILLEGAL_CHARACTERS_RE.sub("", v)) if isinstance(v, str) else v
    state = core.load_state(jid)
    job = core.get_job(jid)
    if not state or not job:
        raise HTTPException(404)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Оформление"
    sty = state.get("style") or {}
    ws.append(["Нарушение", "Подробности", "Страницы", "Доля текста, %", "Пример текста"])
    for i in sty.get("issues") or []:
        ws.append([_safe(i["title"]), _safe(i.get("detail", "")), ", ".join(map(str, i["pages"])), i.get("share"),
                   _safe(i.get("sample", ""))])
    ws2 = wb.create_sheet("Профиль оформления")
    ws2.append(["Что", "Значение", "Доля текста, %"])
    prof = sty.get("profile") or {}
    for f in prof.get("fonts", []):
        ws2.append(["Гарнитура", _safe(", ".join(f.get("names") or [f["family"]])), f["share"]])
    for f in prof.get("sizes", []):
        ws2.append(["Кегль, пт", f["size"], f["share"]])
    for f in prof.get("colors", []):
        ws2.append(["Цвет текста", f["color"], f["share"]])
    ws3 = wb.create_sheet("Справочник продукции")
    ws3.append(["Стр.", "Замечание", "Подробности", "Артикул", "Как в справочнике", "Строка документа"])
    for i in (state.get("products") or {}).get("issues") or []:
        ws3.append([", ".join(map(str, i.get("pages") or [i["page"]])), _safe(i["title"]), _safe(i["detail"]), _safe(i["article"]), _safe(i.get("suggest", "")),
                    _safe(i.get("text", ""))])
    for w in (ws, ws2, ws3):
        for c in w[1]:
            c.font = Font(bold=True)
        for col, width in zip("ABCDEF", (40, 50, 20, 22, 22, 60)):
            w.column_dimensions[col].width = width
        w.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    name = f"Проверки_оформления_{jid}.xlsx"
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(name)}"})


@app.get("/job/{jid}/zip")
def job_zip(jid: int):
    state = core.load_state(jid)
    if not state:
        raise HTTPException(404)
    hide_docx = not state["verify"].get("ok") and not core.load_settings().get("report_allow_unverified")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for kind, fn in state["verify"].get("files", {}).items():
            if kind == "docx" and hide_docx:             # п. 0.9: непринятый отчет не выдается
                continue
            p = os.path.join(core.job_dir(jid), fn)
            if os.path.exists(p):
                z.write(p, os.path.basename(fn))
    return Response(buf.getvalue(), media_type="application/zip",
                    headers={"Content-Disposition": f"attachment; filename=report_{jid}.zip"})


def _idle_job(jid, request=None):
    job = core.get_job(jid)
    if not job:
        raise HTTPException(404)
    user = getattr(request.state, "user", None) if request else None
    if request and not users.can(user, "review") and job.get("owner") != (user or {}).get("login"):
        raise HTTPException(403, "Удалять и перезапускать чужие задания может корректор или администратор")
    if job["status"] in ("running", "uploading"):
        raise HTTPException(409, "Задание сейчас выполняется — дождитесь окончания проверки")
    return job


@app.post("/job/{jid}/delete")
def job_delete(request: Request, jid: int):
    job = _idle_job(jid, request)
    core.delete_job(jid)
    analytics.decisions_changed()
    users.audit(request, "job_delete", jid, job["name"])
    return RedirectResponse("/jobs", status_code=303)


@app.post("/job/{jid}/restart")
def job_restart(request: Request, jid: int):
    job = _idle_job(jid, request)
    users.audit(request, "job_restart", jid, job["name"])
    d = core.job_dir(jid)
    if not os.path.exists(os.path.join(d, "source.pdf")) and not office.find_source(d):
        raise HTTPException(404, "Исходный файл задания не найден")
    with core.state_lock(jid):
        for fn in os.listdir(d):
            p = os.path.join(d, fn)
            if not fn.startswith("source.") and os.path.isfile(p):
                os.remove(p)
        core.update_job(jid, status="queued", progress=0, step="В очереди", error="")
    analytics.decisions_changed()                      # решения этого задания больше не участвуют в обучении
    core.kick()
    return RedirectResponse(f"/job/{jid}", status_code=303)


# ------------------------------------------------------------------ API заданий
@app.get("/api/health")
def health():
    return {"status": "ok", "version": core.VERSION}


@app.get("/api/jobs")
def api_jobs(request: Request):
    return [{k: v for k, v in j.items() if k != "settings"} for j in core.list_jobs(200, owner=users.own_only(request.state.user))]


@app.post("/api/jobs")
async def api_create(request: Request, file: list[UploadFile] = File(...), profile: str = Form(""), mode: str = Form(""),
                     author: str = Form(""), department: str = Form(""), doc_type: str = Form("")):
    ids = await _create_jobs(request, file, profile, mode, "", author, department, doc_type)
    return {"jobs": ids}


@app.get("/api/jobs/{jid}")
def api_job(jid: int):
    job = core.get_job(jid)
    if not job:
        raise HTTPException(404)
    job.pop("settings", None)
    state = core.load_state(jid)
    if state:
        job["verify"] = {k: state["verify"].get(k) for k in ("ok", "passed", "total", "files")}
        job["findings_list"] = state["findings"]
    return job


@app.get("/api/jobs/{jid}/report")
def api_report(jid: int, format: str = "docx"):
    return job_file(jid, format)


async def json_body(request, default=None):
    try:
        body = await request.json()
    except ValueError:
        if default is not None:
            return default
        raise HTTPException(400, "Ожидается JSON")
    if not isinstance(body, dict):
        raise HTTPException(400, "Ожидается JSON-объект")
    return body


def rebuild_job(jid, state=None, summary=None):
    """Пересобирает DOCX и машинную проверку задания. Вызывать под core.state_lock(jid)."""
    job = core.get_job(jid)
    state = state or core.load_state(jid)
    if not state or not job:
        raise HTTPException(404)
    if summary:
        state["meta"]["summary"] = summary
        state["summary_edited"] = True
    st = core.load_settings()
    try:
        st.update(json.loads(job.get("settings") or "{}"))
    except ValueError:
        pass
    state["verify"] = run.build_report(core.job_dir(jid), state, st, lambda m: None)
    core.save_state(jid, state)
    v = state["verify"]
    core.update_job(jid, findings=v.get("count", 0), verify_ok=1 if v.get("ok") else 0,
                    verify_score=f"{v.get('passed', 0)}/{v.get('total', 0)}")
    return v


def _edit_finding(jid, fid, body):
    with core.state_lock(jid):
        state = core.load_state(jid)
        if not state:
            raise HTTPException(404)
        for f in state["findings"]:
            if f["id"] == fid:
                if body.get("status") in ("new", "accepted", "rejected"):
                    f["status"] = body["status"]
                    f.pop("auto", None)                # решение корректора — настоящее, учитывается в обучении
                for k in ("fix", "good", "note", "type"):
                    if isinstance(body.get(k), str):
                        if k == "type" and body[k] not in TYPES:
                            continue
                        if body[k] != f.get(k) and k in ("fix", "good"):
                            f.setdefault("orig_" + k, f.get(k))
                            f["edited"] = True
                        f[k] = body[k][:2000]
                core.save_state(jid, state)
                return {"ok": True, "finding": f}
    raise HTTPException(404)


@app.post("/api/jobs/{jid}/findings/{fid}")
async def api_finding(jid: int, fid: int, request: Request):
    if not users.can(request.state.user, "review"):
        raise HTTPException(403, "Отмечать замечания может корректор или администратор")
    body = await json_body(request)
    res = await run_in_threadpool(_edit_finding, jid, fid, body)
    analytics.decisions_changed()
    users.audit(request, "finding", f"{jid}/{fid}", {k: v for k, v in body.items() if k in ("status", "fix", "good", "type")})
    return res


def _rebuild(jid, summary):
    with core.state_lock(jid):
        return rebuild_job(jid, summary=summary)


@app.post("/api/jobs/{jid}/rebuild")
async def api_rebuild(jid: int, request: Request):
    if not users.can(request.state.user, "review"):
        raise HTTPException(403, "Пересобирать отчет может корректор или администратор")
    body = await json_body(request, {})
    summary = body.get("summary") if isinstance(body.get("summary"), str) else None
    v = await run_in_threadpool(_rebuild, jid, summary)
    users.audit(request, "rebuild", jid, f"{v.get('passed')}/{v.get('total')}")
    return {"ok": v.get("ok"), "passed": v.get("passed"), "total": v.get("total"), "count": v.get("count"),
            "checks": v.get("checks")}


# ------------------------------------------------------------------ аналитика и тренды
def _q(request):
    q = dict(request.query_params)
    q.pop("owner", None)
    own = users.own_only(getattr(request.state, "user", None))
    if own is not None:
        q["owner"] = own                               # автор видит аналитику только своих документов
    return q


@app.get("/analytics/report.pdf")
async def analytics_report_pdf(request: Request):
    """Отчет для руководства за выбранный период и фильтры (те же параметры, что у страницы аналитики)."""
    from . import report_pdf
    q = _q(request)
    q.setdefault("period", "30")
    pdf = await run_in_threadpool(report_pdf.build, q)
    users.audit(request, "mgmt_report", "скачан", {k: v for k, v in q.items() if isinstance(v, str)})
    return Response(pdf, media_type="application/pdf",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(report_pdf.filename())}"})


@app.get("/analytics", response_class=HTMLResponse)
def analytics_page(request: Request):
    return tpl.TemplateResponse(request, "analytics.html", ctx(request, "analytics"))


@app.get("/trends", response_class=HTMLResponse)
def trends_page(request: Request):
    return tpl.TemplateResponse(request, "trends.html", ctx(request, "trends"))


@app.get("/api/analytics")
def api_analytics(request: Request):
    return JSONResponse(analytics.nan_safe(analytics.analytics(_q(request))))


@app.get("/api/trends")
def api_trends(request: Request):
    return JSONResponse(analytics.nan_safe(analytics.trends(_q(request))))


@app.get("/api/analytics/findings")
def api_explore(request: Request, limit: int = 200, offset: int = 0):
    return JSONResponse(analytics.explore(_q(request), max(1, min(limit, 2000)), max(0, offset)))


def _export(sheets, fmt, base):
    stamp = dt.date.today().isoformat()
    if fmt == "csv":
        return Response(analytics.to_csv(sheets), media_type="text/csv; charset=utf-8",
                        headers={"Content-Disposition": f"attachment; filename={base}_{stamp}.csv"})
    return Response(analytics.to_xlsx(sheets, base), headers={"Content-Disposition": f"attachment; filename={base}_{stamp}.xlsx"},
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.get("/api/analytics/export")
def api_analytics_export(request: Request, format: str = "xlsx"):
    q = _q(request)
    if format == "json":
        return JSONResponse(analytics.nan_safe(analytics.analytics(q)),
                            headers={"Content-Disposition": "attachment; filename=korrektor_analytics.json"})
    if format == "findings":
        ex = analytics.explore(q, 10 ** 6)
        head = ["Задание", "Документ", "Дата", "Стр.", "Тип", "Правило", "Источник", "Ошибка", "Как правильно",
                "Статус", "Системно, мест"]
        rows = [[r["job"], r["name"], r["created"], r["page"], r["type"], r["rule_label"], r["source_label"], r["bad"],
                 r["good"], analytics.STATUS.get(r["status"], r["status"]), r["systemic"] or ""] for r in ex["rows"]]
        return _export([("Замечания", head, rows)], q.get("as", "xlsx"), "korrektor_findings")
    return _export(analytics.export_rows(analytics.analytics(q)), format, "korrektor_analytics")


@app.get("/api/trends/export")
def api_trends_export(request: Request, format: str = "xlsx"):
    t = analytics.trends(_q(request))
    if format == "json":
        return JSONResponse(analytics.nan_safe(t), headers={"Content-Disposition": "attachment; filename=korrektor_trends.json"})
    return _export(analytics.trend_rows(t), format, "korrektor_trends")


@app.post("/api/analytics/whitelist")
async def api_whitelist(request: Request):
    """Добавить слова в белый список словаря (из таблицы «Слова вне словаря»)."""
    body = await json_body(request)
    words = body.get("words") if isinstance(body.get("words"), list) else []
    words = [w.strip() for w in words[:500] if isinstance(w, str) and re.fullmatch(r"[\w\-]{2,40}", w.strip())]
    added = []

    def upd(st):
        cur = (st.get("whitelist") or "").split()
        low = {w.lower() for w in cur}
        added[:] = [w for w in dict.fromkeys(words) if w.lower() not in low]
        st["whitelist"] = " ".join(cur + added)
    core.update_settings(upd)
    users.audit(request, "whitelist", ", ".join(added))
    return {"ok": True, "added": added}


@app.post("/api/analytics/glossary")
async def api_glossary(request: Request):
    """Добавить пару «правильно: неправильно» в глоссарий компании (из таблицы «Повторяющиеся ошибки»)."""
    body = await json_body(request)
    good, bad = str(body.get("good", "")).strip(), str(body.get("bad", "")).strip()
    if not good or not bad or good == bad or any(c in good + bad for c in ":\n,") or len(good + bad) > 120:
        raise HTTPException(400, "Нужны непустые «правильно» и «неправильно» без двоеточий, запятых и переводов строк")
    res = {"added": True}

    def upd(st):
        lines = (st.get("glossary") or "").splitlines()
        for i, ln in enumerate(lines):
            if ":" in ln and ln.split(":", 1)[0].strip().lower() == good.lower():
                wrong = [w.strip() for w in ln.split(":", 1)[1].split(",") if w.strip()]
                if bad.lower() in (w.lower() for w in wrong):
                    res["added"] = False
                    return
                lines[i] = f"{ln.split(':', 1)[0].strip()}: " + ", ".join(wrong + [bad])
                break
        else:
            lines.append(f"{good}: {bad}")
        st["glossary"] = "\n".join(lines)
    core.update_settings(upd)
    if res["added"]:
        users.audit(request, "glossary", f"{good}: {bad}")
    return {"ok": True, "added": res["added"]}


@app.post("/api/analytics/bulk")
async def api_bulk(request: Request):
    """Массовая отметка замечаний из обозревателя: [{job, id}] + status; затронутые отчеты пересобираются."""
    body = await json_body(request)
    status = body.get("status")
    if status not in ("new", "accepted", "rejected"):
        raise HTTPException(400, "status: new | accepted | rejected")
    items = collections.defaultdict(set)
    for it in (body.get("items") if isinstance(body.get("items"), list) else [])[:5000]:
        try:
            items[int(it["job"])].add(int(it["id"]))
        except (KeyError, TypeError, ValueError):
            continue
    user = getattr(request.state, "user", None)
    if users.own_only(user) is not None:                # разделение по отделам: только видимые задания
        items = {j: v for j, v in items.items() if users.job_visible(user, j)}
    done = await run_in_threadpool(_bulk, items, status, body.get("rebuild", True) is not False)
    analytics.decisions_changed()
    users.audit(request, "bulk", ", ".join(map(str, done)), {"status": status, "changed": sum(done.values())})
    return {"ok": True, "changed": done}


def _bulk(items, status, rebuild):
    done = {}
    for jid, ids in items.items():
        with core.state_lock(jid):
            state = core.load_state(jid)
            if not state:
                continue
            n = 0
            for f in state["findings"]:
                if f["id"] in ids and f.get("status") != status:
                    f["status"] = status
                    f.pop("auto", None)
                    n += 1
            if not n:
                continue
            core.save_state(jid, state)
            if rebuild:
                rebuild_job(jid, state)
            done[jid] = n
    return done


# ------------------------------------------------------------------ настройки
def grouped():
    return [(g, t, d, [s for s in S if s["group"] == g]) for g, t, d in GROUPS]


@app.get("/settings", response_class=HTMLResponse)
def settings_page(request: Request, saved: int = 0, applied: str = "", bsaved: str = "", berror: str = "",
                  psaved: str = "", perror: str = ""):
    st = core.load_settings()
    if applied and applied in PROFILES:
        st.update(profile_overrides(applied))
    prompt_info = run_prompt_info()
    from . import backup as _bk
    bstat, blist = _bk.status(), _bk.listing(st)
    return tpl.TemplateResponse(request, "settings.html", ctx(request, "settings", groups=grouped(), vals=st, saved=saved,
                                                              applied=applied, prompt=prompt_info,
                                                              bstat=bstat, blist=blist[:15], bsaved=bsaved[:80],
                                                              berror=berror[:300], psaved=psaved[:20],
                                                              perror=perror[:300], pref=_products_info()))


def _products_info():
    from .pipeline import products
    ref = products.load(products.PATH())
    if not ref:
        return {"items": 0, "sync": products.sync_status()} if products.sync_status() else None
    t = os.path.getmtime(products.PATH())
    return {"file": ref.get("file"), "items": len(ref.get("items") or {}), "units": ref.get("units") or {},
            "time": dt.datetime.fromtimestamp(t).strftime("%d.%m.%Y %H:%M"), "sync": products.sync_status()}


def run_prompt_info():
    pd = os.path.join(core.APP_DIR, "prompts")
    custom = core.CUSTOM_PROMPT
    files = sorted(f for f in os.listdir(pd) if f.endswith(".md") and f != "custom_prompt.md")
    active = custom if os.path.exists(custom) else os.path.join(pd, files[-1])
    first = open(active, encoding="utf-8").readline().strip("# \n")
    return {"name": os.path.basename(active), "title": first, "custom": os.path.exists(custom),
            "size": os.path.getsize(active), "path": active}


@app.post("/settings")
async def settings_save(request: Request):
    form = await request.form()
    st = core.load_settings()
    old = dict(st)
    for s in S:
        k = s["key"]
        if s["type"] == "bool":
            st[k] = k in form
        elif s["type"] == "multi":
            st[k] = form.getlist(k)
        elif s["type"] == "secret":
            v = form.get(k, "")
            if v and set(v) != {"•"}:
                st[k] = v.strip()
            elif form.get(k + "__clear"):
                st[k] = ""
        elif k in form:
            st[k] = core.coerce(k, form.get(k))
    changed = [k for k in SCHEMA if old.get(k) != st.get(k)]
    core.save_settings(st)
    users.audit(request, "settings", ", ".join(changed)[:300],
                {k: ("***" if SCHEMA[k]["type"] == "secret" else st.get(k)) for k in changed[:30]})
    return RedirectResponse("/settings?saved=1", status_code=303)


@app.get("/settings/export")
def settings_export(secrets_too: int = 0):
    st = core.load_settings()
    if not secrets_too:
        for s in S:
            if s["type"] == "secret":
                st[s["key"]] = ""
    return Response(json.dumps(st, ensure_ascii=False, indent=1), media_type="application/json",
                    headers={"Content-Disposition": "attachment; filename=korrektor-settings.json"})


@app.post("/settings/import")
async def settings_import(request: Request, file: UploadFile = File(...)):
    try:
        data = json.loads((await file.read()).decode("utf-8"))
        assert isinstance(data, dict)
    except Exception:  # noqa
        raise HTTPException(400, "Файл не является выгрузкой настроек Корректора (ожидается JSON-объект)")
    st = core.load_settings()
    for k, v in data.items():
        if k in SCHEMA and not (SCHEMA[k]["type"] == "secret" and not v):
            st[k] = v if SCHEMA[k]["type"] in ("multi",) else core.coerce(k, v)
    core.save_settings(st)
    users.audit(request, "settings_import", file.filename or "")
    return RedirectResponse("/settings?saved=1", status_code=303)


@app.post("/settings/reset")
def settings_reset(request: Request):
    st = core.load_settings()
    users.audit(request, "settings_reset")
    keep = {k: st.get(k) for k in SCHEMA if SCHEMA[k]["type"] == "secret"}
    new = dict(DEFAULTS)
    new.update(keep)
    core.save_settings(new)
    return RedirectResponse("/settings?saved=1", status_code=303)


@app.post("/settings/prompt")
async def settings_prompt(request: Request, file: UploadFile = File(None), reset: str = Form("")):
    custom = core.CUSTOM_PROMPT
    if reset:
        if os.path.exists(custom):
            os.remove(custom)
    elif file is not None:
        try:
            text = (await file.read()).decode("utf-8")
        except UnicodeDecodeError:
            raise HTTPException(400, "Мастер-промпт должен быть текстовым файлом в кодировке UTF-8 (.md)")
        if "МАСТЕР-ПРОМПТ" not in text[:300]:
            raise HTTPException(400, "Файл не похож на мастер-промпт: в первой строке нет «МАСТЕР-ПРОМПТ»")
        open(custom, "w", encoding="utf-8").write(text)
    users.audit(request, "prompt", "сброс" if reset else (file.filename if file is not None else ""))
    return RedirectResponse("/settings?saved=1#ai", status_code=303)


@app.get("/settings/prompt/download")
def prompt_download():
    info = run_prompt_info()
    return FileResponse(info["path"], filename=info["name"])


@app.post("/settings/test-ai")
def test_ai():
    from .pipeline import llm
    st = core.load_settings()
    try:
        ans = llm.call(st, "Ты проверяешь связь. Ответь строго JSON {\"findings\": []}.", "Проверка связи", print)
        return {"ok": True, "answer": ans[:300]}
    except Exception as e:  # noqa
        return {"ok": False, "error": str(e)[:500]}


@app.post("/settings/test-lt")
def test_lt():
    import httpx
    st = core.load_settings()
    try:
        r = httpx.post(st["lt_url"].rstrip("/") + "/v2/check", data={"text": "Это тестовый текст текст.",
                                                                    "language": "ru-RU"}, timeout=20)
        r.raise_for_status()
        return {"ok": True, "answer": f"LanguageTool отвечает, найдено замечаний: {len(r.json().get('matches', []))}"}
    except Exception as e:  # noqa
        return {"ok": False, "error": str(e)[:300]}


@app.post("/settings/token")
def new_token(request: Request):
    users.audit(request, "token")
    st = core.load_settings()
    st["api_token"] = secrets.token_hex(24)
    core.save_settings(st)
    return {"token": st["api_token"]}


# ------------------------------------------------------------------ документация
def render_md(name, single=False):
    text = open(os.path.join(DOCS_DIR, name + ".md"), encoding="utf-8").read()
    text = re.sub(r"\]\(images/", "](/docs-files/images/", text)
    if single:
        text = re.sub(r"\]\((\w+)\.md(#[^)]+)?\)", lambda m: f"]({m.group(2) or '#doc-' + m.group(1)})", text)
    else:
        text = re.sub(r"\]\((\w+)\.md(#[^)]+)?\)", lambda m: f"](/docs/{m.group(1)}{m.group(2) or ''})", text)
    html = markdown.markdown(text, extensions=["tables", "fenced_code", "toc", "attr_list", "admonition"],
                             extension_configs={"toc": {"slugify": slugify_unicode}})
    return re.sub(r'<pre><code class="language-mermaid">(.*?)</code></pre>',
                  lambda m: '<pre class="mermaid">' + m.group(1) + "</pre>", html, flags=re.S)


@app.get("/docs/all", response_class=HTMLResponse)
def docs_all(request: Request):
    html = "".join(f'<section id="doc-{n}" style="page-break-before:always">{render_md(n, True)}</section>' for n, _ in DOC_PAGES)
    return tpl.TemplateResponse(request, "docs.html", ctx(request, "docs", html=html, pages=DOC_PAGES, current="all",
                                                          title="Вся документация"))


@app.get("/docs", response_class=HTMLResponse)
@app.get("/docs/{name}", response_class=HTMLResponse)
def docs(request: Request, name: str = "index"):
    names = dict(DOC_PAGES)
    if name not in names:
        raise HTTPException(404)
    return tpl.TemplateResponse(request, "docs.html", ctx(request, "docs", html=render_md(name), pages=DOC_PAGES,
                                                          current=name, title=names[name]))


@app.get("/about", response_class=HTMLResponse)
def about(request: Request):
    return RedirectResponse("/docs/methodology", status_code=303)


admin.setup(app, tpl, ctx)
